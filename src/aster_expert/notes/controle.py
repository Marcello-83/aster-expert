"""Automatic check of the Word notes (lot 45), read back from the `.docx` itself.

For each note:

- every point of the fiche appears (a title « Point N » of the hypotheses note, a row of
  the limits table of the results note);
- every number quoted is identical to its source: the values of the data file, the
  coefficients of the combinations, the cells of the declared tables of the results
  note (re-read from the package), and every other measured value of a results table
  must exist in a file of the package within the digits displayed;
- the decisions of the fiche appear word for word (nothing is rewritten by hand);
- no `[À CONFIRMER`, `[QUESTION` or `TODO` is left;
- the layout is verified from the file: Arial, 10 pt, black text, header rows of the
  tables on a grey background (never blue), numbered titles, no theme colour.

The check reports what it found; it never corrects.
"""

from __future__ import annotations

import bisect
import csv
import json
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from docx import Document
from docx.document import Document as DocumentType
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.table import Table
from docx.text.paragraph import Paragraph

from aster_expert.agent.hypotheses import (
    FindingReference,
    HypothesesSheet,
    PracticeHypothesisEntry,
)
from aster_expert.combinaisons.enumeration import enumerate_combinations
from aster_expert.notes.docx_style import (
    FONT,
    HEADER_FILL,
    SIZE_PT,
    TEXT_COLOR,
    neutralise_markers,
)
from aster_expert.notes.formats import displayed_tolerance, parse_number
from aster_expert.notes.hypotheses import (
    constat_rows,
    point_title,
    value_text,
)
from aster_expert.notes.tables import resolve_table
from aster_expert.projet.context import ProjectContext

FORBIDDEN = ("[À CONFIRMER", "[QUESTION", "TODO")
HEADING_PATTERN = re.compile(r"^\d+(\.\d+)*\.? \S")


@dataclass
class NoteCheck:
    """What the check found in one note."""

    note: str
    counts: dict[str, int] = field(default_factory=dict)
    problems: list[str] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return not self.problems

    def count(self, key: str, n: int = 1) -> None:
        self.counts[key] = self.counts.get(key, 0) + n

    def problem(self, message: str) -> None:
        self.problems.append(message)


# -- reading the .docx back ----------------------------------------------------------


@dataclass
class Block:
    kind: str  # "p" or "t"
    paragraph: Paragraph | None = None
    table: Table | None = None


def blocks(document: DocumentType) -> list[Block]:
    out: list[Block] = []
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            out.append(Block("p", paragraph=Paragraph(child, document)))
        elif child.tag == qn("w:tbl"):
            out.append(Block("t", table=Table(child, document)))
    return out


def all_paragraphs(document: DocumentType) -> list[Paragraph]:
    found = list(document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                found.extend(cell.paragraphs)
    for section in document.sections:
        found.extend(section.footer.paragraphs)
        found.extend(section.header.paragraphs)
    return found


def full_text(document: DocumentType) -> str:
    return "\n".join(p.text for p in all_paragraphs(document))


def table_rows(table: Table) -> list[list[str]]:
    return [[cell.text for cell in row.cells] for row in table.rows]


def tables_by_caption(document: DocumentType) -> dict[str, list[Table]]:
    """Each table under the last non-empty paragraph that precedes it."""
    found: dict[str, list[Table]] = {}
    caption = ""
    for block in blocks(document):
        if block.kind == "p" and block.paragraph is not None:
            if block.paragraph.text.strip():
                caption = block.paragraph.text.strip()
        elif block.table is not None:
            found.setdefault(caption, []).append(block.table)
    return found


def _norm(text: str) -> str:
    """Whitespace-normalised, as the note shows it (markers quoted by the fiche are
    shown between guillemets)."""
    return " ".join(neutralise_markers(text).split())


# -- layout and markers --------------------------------------------------------------


def check_layout(document: DocumentType, path: Path, result: NoteCheck) -> None:
    runs = 0
    for paragraph in all_paragraphs(document):
        for run in paragraph.runs:
            if run._r.find(qn("w:drawing")) is not None:
                continue  # a figure, not text
            runs += 1
            if run.font.name != FONT:
                result.problem(
                    f"police {run.font.name!r} au lieu de {FONT} : {run.text[:40]!r}"
                )
            if run.font.size != Pt(SIZE_PT):
                result.problem(
                    f"taille {run.font.size} au lieu de {SIZE_PT} pt : {run.text[:40]!r}"
                )
            color = run.font.color.rgb if run.font.color is not None else None
            if str(color) != TEXT_COLOR:
                result.problem(f"couleur {color} au lieu de noir : {run.text[:40]!r}")
    result.count("runs_verifies", runs)
    styles_used = {
        p.style.name for p in all_paragraphs(document) if p.style is not None
    }
    for name in sorted(styles_used):
        style = document.styles[name]
        color = style.font.color.rgb if style.font.color is not None else None
        if (
            style.font.name != FONT
            or style.font.size != Pt(SIZE_PT)
            or str(color) != TEXT_COLOR
        ):
            result.problem(
                f"style « {name} » : {style.font.name}, {style.font.size}, {color}"
            )
    shaded = 0
    for table in document.tables:
        header = table.rows[0]
        for cell in header.cells:
            fills = [s.get(qn("w:fill")) for s in cell._tc.iter(qn("w:shd"))]
            if fills != [HEADER_FILL]:
                result.problem(
                    f"en-tête de tableau sans fond gris {HEADER_FILL} : {cell.text[:30]!r}"
                )
            if not all(r.font.bold for p in cell.paragraphs for r in p.runs):
                result.problem(f"en-tête de tableau hors gras : {cell.text[:30]!r}")
            shaded += 1
        for row in table.rows[1:]:
            for cell in row.cells:
                if any(True for _ in cell._tc.iter(qn("w:shd"))):
                    result.problem(f"cellule de corps avec fond : {cell.text[:30]!r}")
    result.count("cellules_d_en_tete_grises", shaded)
    result.count("tableaux", len(document.tables))
    with zipfile.ZipFile(path) as archive:
        for part in ("word/document.xml", "word/styles.xml", "word/footer1.xml"):
            if part in archive.namelist():
                xml = archive.read(part).decode("utf-8")
                for marker in ("themeColor", "w:highlight"):
                    if marker in xml:
                        result.problem(f"{part} contient {marker}")
                for fill in set(re.findall(r'w:fill="([0-9A-Fa-f]{6})"', xml)):
                    if fill.upper() != HEADER_FILL:
                        result.problem(
                            f"{part} : fond {fill} autre que le gris {HEADER_FILL}"
                        )
    headings = [
        p
        for p in document.paragraphs
        if p.style is not None and p.style.name.startswith("Heading")
    ]
    expected = [0, 0, 0]
    for heading in headings:
        text = heading.text
        if not HEADING_PATTERN.match(text):
            result.problem(f"titre non numéroté : {text!r}")
            continue
        number = text.split()[0].rstrip(".")
        parts = [int(x) for x in number.split(".")]
        level = len(parts)
        expected[level - 1] += 1
        for deeper in range(level, 3):
            expected[deeper] = 0
        if parts != expected[:level]:
            result.problem(
                f"numérotation incohérente : {text!r} (attendu {expected[:level]})"
            )
    result.count("titres_numerotes", len(headings))


ANNEX_TITLE = "Annexe : historique des révisions"
_QUOTED_MARKER = re.compile(r"«\s*(?:À CONFIRMER|QUESTION)[^»]*»")
_OPEN_MARKERS = ("[À CONFIRMER", "[QUESTION")
_SOURCE_MARKER = re.compile(r"\[(?:À CONFIRMER|QUESTION)[^\]]*\]|TODO")


def _annex_start(document: DocumentType) -> int | None:
    """Index, among the paragraphs of the body, of the annex heading (if any)."""
    for index, paragraph in enumerate(document.paragraphs):
        if paragraph.text.strip().endswith(ANNEX_TITLE):
            return index
    return None


def check_markers(document: DocumentType, result: NoteCheck) -> None:
    """Placeholder markers (lot 45, refined at lot 46).

    A marker in a text in force fails the check, whatever its form (`[À CONFIRMER MM]`
    or quoted `« À CONFIRMER MM »`). A marker in the history (the annex of revisions:
    decisions replaced, revision events) is allowed when it is quoted between
    guillemets, as the note shows a marker cited from a closed revision; an open
    bracketed marker fails even there. `TODO` fails everywhere.
    """
    # Table cells carry most of the quoted history: `blocks` orders paragraphs and
    # tables as the body does, the annex starts at its heading.
    annex = _annex_start(document)
    in_force_text: list[str] = []
    history_text: list[str] = []
    in_annex = False
    for block in blocks(document):
        if block.kind == "p" and block.paragraph is not None:
            if annex is not None and block.paragraph.text.strip().endswith(ANNEX_TITLE):
                in_annex = True
            (history_text if in_annex else in_force_text).append(block.paragraph.text)
        elif block.table is not None:
            cells = [c.text for row in block.table.rows for c in row.cells]
            (history_text if in_annex else in_force_text).extend(cells)
    for section in document.sections:
        in_force_text.extend(p.text for p in section.footer.paragraphs)
        in_force_text.extend(p.text for p in section.header.paragraphs)

    current = "\n".join(in_force_text)
    past = "\n".join(history_text)
    for marker in (*_OPEN_MARKERS, "TODO"):
        for where, text in (("texte en vigueur", current), ("historique", past)):
            found = text.count(marker)
            if found:
                result.problem(f"{found} occurrence(s) de « {marker} » dans le {where}")
    quoted_in_force = len(_QUOTED_MARKER.findall(current))
    if quoted_in_force:
        result.problem(
            f"{quoted_in_force} marqueur(s) cité(s) entre guillemets dans un texte en "
            "vigueur : seul l'historique des révisions peut en citer"
        )
    result.count("marqueurs_cites_historique", len(_QUOTED_MARKER.findall(past)))
    result.count("marqueurs_interdits_recherches", len(_OPEN_MARKERS) + 1)


def check_sheet_markers(sheet: HypothesesSheet, result: NoteCheck) -> None:
    """A marker in a decision in force, in the fiche itself (the note would quote it)."""
    for entry in sheet.entries:
        if not isinstance(entry, PracticeHypothesisEntry):
            continue
        for name in ("option_retained", "justification", "validity_domain"):
            value = getattr(entry, name)
            if value and _SOURCE_MARKER.search(value):
                result.problem(
                    f"point {entry.point_label} : marqueur dans {name}, texte en vigueur"
                )


# -- hypotheses note -----------------------------------------------------------------


def _expected_points(sheet: HypothesesSheet) -> set[str | None]:
    points: set[str | None] = {e.point_label for e in sheet.entries}
    points |= {g.point_label for g in sheet.group_naming}
    points |= {c.point_label for c in sheet.planned_commands}
    return points


def _cell_number_matches(text: str, expected: float) -> bool:
    got = parse_number(text)
    return got is not None and got == expected


def _check_data(document: DocumentType, ctx: ProjectContext, result: NoteCheck) -> None:
    data = ctx.data()
    tables = tables_by_caption(document)
    for group, values in data.groups.items():
        found = tables.get(f"Données — {group}")
        if not found:
            result.problem(f"données : tableau du groupe {group} absent")
            continue
        rows = {r[0]: r for r in table_rows(found[0])[1:]}
        for key, value in values.items():
            result.count("valeurs_de_donnees_verifiees")
            row = rows.get(key)
            if row is None:
                result.problem(f"données : {group}.{key} absente du tableau")
                continue
            if row[2] != value.unit:
                result.problem(
                    f"données : {group}.{key} unité {row[2]!r} au lieu de {value.unit!r}"
                )
            if value.status == "a_fournir":
                continue
            if value.profile is not None:
                continue
            raw = value.value
            if isinstance(raw, tuple):
                shown = [parse_number(x) for x in row[1].split(" ; ")]
                if shown != [float(x) for x in raw]:
                    result.problem(
                        f"données : {group}.{key} = {raw!r}, note {row[1]!r}"
                    )
            elif isinstance(raw, int | float):
                if not _cell_number_matches(row[1], float(raw)):
                    result.problem(
                        f"données : {group}.{key} = {raw!r}, note {row[1]!r}"
                    )
            elif row[1] != str(raw):
                result.problem(f"données : {group}.{key} = {raw!r}, note {row[1]!r}")
            if value_text(value) != row[1]:
                result.problem(
                    f"données : {group}.{key} : texte {row[1]!r} différent de la source"
                )
        for key, value in values.items():
            if value.profile is None:
                continue
            profile = tables.get(f"Profil {group}.{key} (points z, p)")
            if not profile:
                result.problem(f"données : profil {group}.{key} absent")
                continue
            lines = table_rows(profile[0])[1:]
            if len(lines) != len(value.profile):
                result.problem(
                    f"profil {group}.{key} : {len(lines)} lignes pour {len(value.profile)} points"
                )
                continue
            for point, line in zip(value.profile, lines, strict=True):
                result.count("valeurs_de_donnees_verifiees", 2)
                if not (
                    _cell_number_matches(line[0], point.z)
                    and _cell_number_matches(line[1], point.p)
                ):
                    result.problem(
                        f"profil {group}.{key} : point ({point.z}, {point.p}) différent de la note"
                    )


def _check_combinations(
    document: DocumentType, ctx: ProjectContext, result: NoteCheck
) -> None:
    section = ctx.declaration.combinations
    if section is None:
        return
    combos = enumerate_combinations(section, ctx.number)
    found = next(
        (
            t
            for c, t in tables_by_caption(document).items()
            if c.startswith("Combinaisons (")
        ),
        None,
    )
    if not found:
        result.problem("combinaisons : tableau absent")
        return
    rows = {r[0]: r for r in table_rows(found[0])[1:]}
    for combo in combos:
        result.count("combinaisons_verifiees")
        row = rows.get(combo.name)
        if row is None:
            result.problem(f"combinaison {combo.name} absente")
            continue
        shown: dict[str, float | None] = {}
        for part in row[2].split(" ; "):
            case, _, coefficient = part.partition(" × ")
            shown[case] = parse_number(coefficient)
        if shown != {k: float(v) for k, v in combo.coefficients.items()}:
            result.problem(
                f"combinaison {combo.name} : coefficients {row[2]!r} différents de {combo.coefficients}"
            )


def check_hypotheses_note(ctx: ProjectContext, path: Path) -> NoteCheck:
    result = NoteCheck(note=path.name)
    document = Document(str(path))
    sheet = HypothesesSheet.model_validate_json(
        ctx.fiche_path.read_text(encoding="utf-8")
    )
    text = _norm(full_text(document))
    heading_texts = [
        p.text
        for p in document.paragraphs
        if p.style is not None and p.style.name.startswith("Heading")
    ]

    # every point appears under its title
    for label in sorted(_expected_points(sheet), key=lambda x: (x is None, str(x))):
        result.count("points_verifies")
        title = point_title(label)
        if not any(h.endswith(title) for h in heading_texts):
            result.problem(f"{title} : aucun titre")

    # decisions word for word
    for entry in sheet.entries:
        if (
            not isinstance(entry, PracticeHypothesisEntry)
            or entry.status == "non_tranche"
        ):
            continue
        for name in ("option_retained", "justification", "validity_domain"):
            value = getattr(entry, name)
            if value:
                result.count("textes_de_decision_verifies")
                for part in value.split("\n"):
                    if part.strip() and _norm(part) not in text:
                        result.problem(
                            f"point {entry.point_label} : {name} absent mot pour mot ({_norm(part)[:60]!r})"
                        )
        for variant in entry.discarded_variants:
            result.count("textes_de_decision_verifies")
            if (
                _norm(variant.description) not in text
                or _norm(variant.reason) not in text
            ):
                result.problem(
                    f"point {entry.point_label} : variante écartée absente ({variant.description[:50]!r})"
                )
        for source in entry.sources:
            result.count("sources_verifiees")
            if _norm(source.reference) not in text:
                result.problem(
                    f"point {entry.point_label} : source absente ({source.reference[:50]!r})"
                )
        if entry.decided_by is None or entry.decided_at is None:
            result.problem(
                f"point {entry.point_label} : décideur ou date manquant dans la fiche"
            )
        else:
            result.count("decideurs_et_dates_verifies")
            if (
                entry.decided_at.date().isoformat() not in text
                or entry.decided_by not in text
            ):
                result.problem(
                    f"point {entry.point_label} : décideur ou date absents de la note"
                )
    for entry in sheet.entries:
        if isinstance(entry, PracticeHypothesisEntry):
            for prior in entry.decision_history:
                result.count("decisions_remplacees_verifiees")
                if _norm(prior.option_retained)[:80] not in text:
                    result.problem(
                        f"point {entry.point_label} : décision remplacée absente de l'annexe"
                    )

    # data, naming convention, findings, combinations
    _check_data(document, ctx, result)
    for group in sheet.group_naming:
        result.count("groupes_nommes_verifies")
        if group.group_name not in text:
            result.problem(f"groupe {group.group_name} absent")
    rows, _ = constat_rows(sheet, ctx.repo_root)
    cited = [e.constat_id for e in sheet.entries if isinstance(e, FindingReference)]
    table = next(
        (
            t
            for c, t in tables_by_caption(document).items()
            if c.startswith("Constats cités (")
        ),
        None,
    )
    shown = {r[0]: r for r in table_rows(table[0])[1:]} if table else {}
    for constat in cited:
        result.count("constats_verifies")
        row = shown.get(constat)
        base_row = next((r for r in rows if r.constat_id == constat), None)
        if row is None or base_row is None:
            result.problem(f"constat {constat} absent du tableau")
        elif (row[4], row[5]) != (base_row.level, base_row.status):
            result.problem(
                f"constat {constat} : niveau/statut {row[4], row[5]} différents "
                f"de {base_row.level, base_row.status}"
            )
    _check_combinations(document, ctx, result)
    check_markers(document, result)
    check_sheet_markers(sheet, result)
    check_layout(document, path, result)
    return result


# -- results note --------------------------------------------------------------------


def _numbers_of_package(directory: Path) -> list[float]:
    numbers: list[float] = []

    def walk(value: Any) -> None:
        if isinstance(value, bool):
            return
        if isinstance(value, int | float):
            numbers.append(float(value))
        elif isinstance(value, dict):
            for item in value.values():
                walk(item)
        elif isinstance(value, list):
            for item in value:
                walk(item)

    for path in directory.rglob("*.json"):
        walk(json.loads(path.read_text(encoding="utf-8")))
    for path in directory.rglob("*.csv"):
        with path.open(encoding="utf-8", newline="") as stream:
            for row in csv.reader(stream):
                for text in row:
                    try:
                        numbers.append(float(text))
                    except ValueError:
                        continue
    numbers.sort()
    return numbers


def _in_package(sorted_numbers: list[float], text: str) -> bool:
    value = parse_number(text)
    if value is None:
        return True
    tolerance = displayed_tolerance(text)
    window = tolerance * abs(value) + 1e-300
    low = bisect.bisect_left(sorted_numbers, value - window)
    return low < len(sorted_numbers) and sorted_numbers[low] <= value + window


def check_results_note(ctx: ProjectContext, path: Path) -> NoteCheck:
    from aster_expert.notes.resultats import Package  # noqa: PLC0415

    result = NoteCheck(note=path.name)
    document = Document(str(path))
    package = Package(ctx)
    sheet = HypothesesSheet.model_validate_json(
        ctx.fiche_path.read_text(encoding="utf-8")
    )
    tables = tables_by_caption(document)
    text = _norm(full_text(document))

    # every point of the fiche appears (limits table)
    limits = next(
        (t for c, t in tables.items() if "décisions tranchées de la fiche" in c), None
    )
    labels = {r[0] for r in table_rows(limits[0])[1:]} if limits else set()
    for label in sorted(_expected_points(sheet), key=lambda x: (x is None, str(x))):
        result.count("points_verifies")
        if (label if label is not None else "—") not in labels:
            result.problem(f"{point_title(label)} : absent du tableau des limites")

    # declared tables, cell by cell against the package
    for declared in package.declared.note_tables:
        found = tables.get(declared.title)
        if not found:
            result.problem(f"tableau « {declared.title} » absent")
            continue
        resolved = resolve_table(package.directory, declared)
        shown = table_rows(found[0])
        if shown[0] != resolved.headers:
            result.problem(f"tableau « {declared.title} » : en-têtes {shown[0]}")
        if len(shown) - 1 != resolved.shown_rows:
            result.problem(
                f"tableau « {declared.title} » : {len(shown) - 1} lignes pour {resolved.shown_rows}"
            )
            continue
        for expected_row, values, line in zip(
            resolved.cells, resolved.values, shown[1:], strict=True
        ):
            for want, raw, got in zip(expected_row, values, line, strict=True):
                result.count("cellules_de_tableau_verifiees")
                if got != want:
                    result.problem(
                        f"tableau « {declared.title} » : {got!r} au lieu de {want!r}"
                    )
                elif isinstance(raw, float) and not _tolerates(got, raw):
                    result.problem(
                        f"tableau « {declared.title} » : {got!r} s'écarte de la source {raw!r}"
                    )
    result.count("figures_declarees", len(package.declared.note_figures))
    result.count("figures_dans_la_note", len(document.inline_shapes))

    # every measured value of any table exists in the package
    sorted_numbers = _numbers_of_package(package.directory)
    for table in document.tables:
        for row in table_rows(table)[1:]:
            for cell in row:
                if "," in cell or "e" in cell.lower():
                    value = parse_number(cell)
                    if value is None:
                        continue
                    result.count("valeurs_mesurees_retrouvees_dans_le_paquet")
                    if not _in_package(sorted_numbers, cell):
                        result.problem(
                            f"valeur {cell!r} introuvable dans les fichiers du paquet"
                        )

    # level and criteria
    level = package.manifest["niveau"]["atteint"] or "aucun"
    result.count("niveau_verifie")
    if f"Niveau atteint : {level}" not in text:
        result.problem(f"niveau atteint {level} absent de la note")
    for criterion in package.manifest["criteres"]:
        result.count("criteres_verifies")
        if criterion["document"] not in text:
            result.problem(f"critère {criterion['document']} absent")
        if (
            not criterion.get("anteriorite_demontrable", False)
            and "antériorité non démontrable" not in text
        ):
            result.problem("antériorité non démontrable : mention absente")
    check_markers(document, result)
    check_layout(document, path, result)
    return result


def _tolerates(text: str, source: float) -> bool:
    shown = parse_number(text)
    if shown is None:
        return False
    return abs(shown - source) <= displayed_tolerance(text) * abs(source) + 1e-300


def sheet_points(sheet: HypothesesSheet) -> list[str | None]:
    return sorted(_expected_points(sheet), key=lambda x: (x is None, str(x)))


__all__ = [
    "FORBIDDEN",
    "NoteCheck",
    "check_hypotheses_note",
    "check_layout",
    "check_markers",
    "check_sheet_markers",
    "check_results_note",
]
