"""Common layout of the Word notes (lot 45), wanted by MM: Arial 10, black text,
table headers on a grey background (never blue), numbered titles.

Every run carries its font, size and colour explicitly (not only through the
styles), so that `aster_expert.notes.controle` can verify the layout by reading the
`.docx` back. The numbering of the titles is written in the text (`1.`, `1.1`):
it does not depend on a Word numbering definition.

Nothing here knows an ouvrage: a note is a succession of titles, paragraphs,
tables and figures, each filled by the caller from the fiche, the data or the
package of results.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

FONT = "Arial"
SIZE_PT = 10
TEXT_COLOR = "000000"
HEADER_FILL = "D9D9D9"
"""Grey of the table headers."""
PAGE_WIDTH_CM = 21.0
MARGIN_CM = 2.0
TEXT_WIDTH_CM = PAGE_WIDTH_CM - 2 * MARGIN_CM
HEADING_STYLES = {1: "Heading 1", 2: "Heading 2", 3: "Heading 3"}
_MARKER = re.compile(r"\[(À CONFIRMER|QUESTION)([^\]]*)\]")


def neutralise_markers(text: str) -> str:
    """A placeholder marker quoted by a historical text of the fiche
    (`[À CONFIRMER MM]`) is shown between guillemets (`« À CONFIRMER MM »`): a note
    never carries an open marker, and the quotation stays readable."""
    return _MARKER.sub(lambda m: f"« {m.group(1)}{m.group(2)} »", text)


def _set_fonts(rpr: Any) -> None:
    """`w:rFonts` with the four slots of Word, so that no script falls back to
    another font."""
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attribute in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attribute), FONT)
    for theme in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if fonts.get(qn(theme)) is not None:
            del fonts.attrib[qn(theme)]


def _style_run_properties(
    run: Any, *, bold: bool = False, italic: bool = False
) -> None:
    run.font.name = FONT
    run.font.size = Pt(SIZE_PT)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._r.get_or_add_rPr()
    _set_fonts(rpr)
    color = rpr.find(qn("w:color"))
    if color is not None and color.get(qn("w:themeColor")) is not None:
        del color.attrib[qn("w:themeColor")]


def _configure_style(style: Any, *, bold: bool) -> None:
    style.font.name = FONT
    style.font.size = Pt(SIZE_PT)
    style.font.bold = bold
    style.font.italic = False
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    _set_fonts(rpr)
    color = rpr.find(qn("w:color"))
    if color is not None:
        for attribute in ("w:themeColor", "w:themeShade", "w:themeTint"):
            if color.get(qn(attribute)) is not None:
                del color.attrib[qn(attribute)]


def _shade(cell: _Cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    for existing in properties.findall(qn("w:shd")):
        properties.remove(existing)
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _repeat_header(row: Any) -> None:
    properties = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    properties.append(marker)


def _no_split(row: Any) -> None:
    properties = row._tr.get_or_add_trPr()
    properties.append(OxmlElement("w:cantSplit"))


def _field(paragraph: Paragraph, instruction: str) -> None:
    """A Word field (`PAGE`, `NUMPAGES`) in a run formatted like the text."""
    for kind in ("begin", None, "end"):
        run = paragraph.add_run()
        _style_run_properties(run)
        if kind is None:
            instr = OxmlElement("w:instrText")
            instr.set(qn("xml:space"), "preserve")
            instr.text = f" {instruction} "
            run._r.append(instr)
        else:
            mark = OxmlElement("w:fldChar")
            mark.set(qn("w:fldCharType"), kind)
            run._r.append(mark)


class NoteDocument:
    """A note under construction."""

    def __init__(self, *, title: str, footer: str) -> None:
        self.document: DocumentType = Document()
        self._counters = [0, 0, 0]
        self._title = title
        self.tables = 0
        self.figures = 0
        self._configure(footer)

    # -- page and styles -------------------------------------------------------

    def _configure(self, footer: str) -> None:
        document = self.document
        section = document.sections[0]
        section.page_width = Cm(PAGE_WIDTH_CM)
        section.page_height = Cm(29.7)
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(section, side, Cm(MARGIN_CM))
        self._prune_styles()
        for style in document.styles:
            if getattr(style, "type", None) in (1, 2) and hasattr(style, "font"):
                # paragraph (1) and character (2) styles: Arial 10 black everywhere
                _configure_style(style, bold=style.name in HEADING_STYLES.values())
        normal = document.styles["Normal"]
        normal.paragraph_format.space_after = Pt(4)
        normal.paragraph_format.space_before = Pt(0)
        for style_name in HEADING_STYLES.values():
            heading = document.styles[style_name]
            heading.paragraph_format.space_before = Pt(12)
            heading.paragraph_format.space_after = Pt(4)
            heading.paragraph_format.keep_with_next = True
        properties = document.core_properties
        properties.title = self._title
        properties.author = ""
        properties.comments = ""
        paragraph = section.footer.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(f"{footer} -- page ")
        _style_run_properties(run)
        _field(paragraph, "PAGE")
        run = paragraph.add_run(" sur ")
        _style_run_properties(run)
        _field(paragraph, "NUMPAGES")

    # -- text ------------------------------------------------------------------

    def _prune_styles(self) -> None:
        """Remove the styles of the default template that the note does not use:
        most carry theme colours (blue accents), none must be able to leak."""
        keep = {
            "Normal",
            "Header",
            "Footer",
            "Table Grid",
            "Normal Table",
            "Default Paragraph Font",
            "No List",
            *HEADING_STYLES.values(),
        }
        for style in list(self.document.styles):
            if style.name not in keep:
                style.element.getparent().remove(style.element)
        for element in self.document.styles.element.iter():
            for attribute in ("w:themeColor", "w:themeShade", "w:themeTint"):
                if element.get(qn(attribute)) is not None:
                    del element.attrib[qn(attribute)]

    def title(self, text: str, subtitle: str | None = None) -> None:
        paragraph = self.document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(text.upper())
        _style_run_properties(run, bold=True)
        if subtitle:
            sub = self.document.add_paragraph()
            sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _style_run_properties(sub.add_run(subtitle))

    def heading(self, text: str, level: int = 1) -> str:
        """A numbered title; returns its number (`2.3`)."""
        if level not in HEADING_STYLES:
            message = f"niveau de titre {level} non géré (1 à 3)"
            raise ValueError(message)
        self._counters[level - 1] += 1
        for deeper in range(level, 3):
            self._counters[deeper] = 0
        number = ".".join(str(n) for n in self._counters[:level])
        paragraph = self.document.add_paragraph(style=HEADING_STYLES[level])
        run = paragraph.add_run(
            f"{number}{'.' if level == 1 else ''} {neutralise_markers(text)}"
        )
        _style_run_properties(run, bold=True)
        return number

    def paragraph(
        self,
        text: str = "",
        *,
        bold: bool = False,
        italic: bool = False,
        label: str | None = None,
    ) -> Paragraph:
        """A paragraph; `label` is written in bold before `text`."""
        paragraph = self.document.add_paragraph()
        text = neutralise_markers(text)
        if label:
            _style_run_properties(paragraph.add_run(f"{label} : "), bold=True)
        if text or not label:
            _style_run_properties(paragraph.add_run(text), bold=bold, italic=italic)
        return paragraph

    def lines(self, label: str, items: list[str]) -> None:
        """A label followed by one paragraph per item, each introduced by a dash."""
        self.paragraph(label, bold=True)
        for item in items:
            paragraph = self.document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.6)
            _style_run_properties(paragraph.add_run(f"– {neutralise_markers(item)}"))

    # -- tables ----------------------------------------------------------------

    def table(
        self,
        headers: list[str],
        rows: list[list[str]],
        *,
        widths_cm: list[float] | None = None,
        caption: str | None = None,
    ) -> Table:
        """A table with its header on a grey background, repeated on each page."""
        if caption:
            self.paragraph(caption, bold=True)
        columns = len(headers)
        table = self.document.add_table(rows=1, cols=columns)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        if widths_cm is None:
            widths_cm = [TEXT_WIDTH_CM / columns] * columns
        scale = TEXT_WIDTH_CM / sum(widths_cm)
        widths = [Cm(w * scale) for w in widths_cm]
        header = table.rows[0]
        _repeat_header(header)
        for index, text in enumerate(headers):
            cell = header.cells[index]
            _shade(cell, HEADER_FILL)
            self._fill(cell, text, bold=True)
        for values in rows:
            row = table.add_row()
            _no_split(row)
            for index in range(columns):
                self._fill(
                    row.cells[index], values[index] if index < len(values) else ""
                )
        for index, width in enumerate(widths):
            table.columns[index].width = width
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = width
        self.document.add_paragraph().paragraph_format.space_after = Pt(2)
        self.tables += 1
        assert isinstance(table, Table)
        return table

    def _fill(self, cell: _Cell, text: str, *, bold: bool = False) -> None:
        parts = neutralise_markers(text).split("\n")
        paragraph = cell.paragraphs[0]
        for index, part in enumerate(parts):
            if index:
                paragraph = cell.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(1)
            _style_run_properties(paragraph.add_run(part), bold=bold)

    def key_values(
        self, pairs: list[tuple[str, str]], *, headers: tuple[str, str]
    ) -> Table:
        return self.table(
            list(headers), [[k, v] for k, v in pairs], widths_cm=[5.0, 12.0]
        )

    # -- figures ---------------------------------------------------------------

    def figure(self, path: Path, caption: str, *, width_cm: float = 16.0) -> None:
        paragraph = self.document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.keep_with_next = True
        paragraph.add_run().add_picture(
            str(path), width=Cm(min(width_cm, TEXT_WIDTH_CM))
        )
        legend = self.document.add_paragraph()
        legend.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.figures += 1
        _style_run_properties(
            legend.add_run(f"Figure {self.figures} : {caption}"), italic=True
        )

    def page_break(self) -> None:
        self.document.add_page_break()  # type: ignore[no-untyped-call]

    def save(self, path: Path) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        self.document.save(str(path))


__all__ = [
    "FONT",
    "HEADER_FILL",
    "SIZE_PT",
    "TEXT_COLOR",
    "TEXT_WIDTH_CM",
    "NoteDocument",
    "neutralise_markers",
]
