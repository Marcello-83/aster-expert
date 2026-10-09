"""Summary workbook of the deliverables (lots 49 and 49 bis), built from the results
package only.

Nothing is recomputed: every value is read in a file of `resultats/paquet/` (section
`visualisation`: the engineering summary of lot 49 bis and the files of lot 49; the
sections `controles`, `niveau_e`, `contre_calcul`; the manifest of the package).
Layout asked by MM: font Arial 10, black text, table headers on a grey fill, never
blue; frozen headers, adapted column widths, number formats from the display unit
(`config/arrondis.yaml`: the cell holds the exact converted value, only its display is
rounded). A value read where the model has no physical meaning (« non significative »)
sits on a light grey fill, with its reason in the row.

Sheets (lot 49 bis): « Synthèse ingénieur » (one block per family of members, one row
per section and design quantity, per limit state the value and its concomitant N, M,
V), « Vérifications de service », « Torseurs » (wide: one row per section × limit
state × quantity, then one per section × combination), « Détail <famille> » (one
member per row and limit state), then the sheets of lot 49 kept: reactions and
equilibrium, controls D and E, external comparison, conventions, provenance. The long
raw tables of lot 49 (values at reading points, envelopes along the members, extrema
per field) are CSV annexes, out of the workbook.

Lot 49 ter: the « Synthèse ingénieur » is in the engineer's convention of the CSV
files (N > 0 in compression, M signed as the section declares: span or support moment,
|V|) with physical names (D4); « Torseurs » and the raw columns keep the Code_Aster
convention, and the head of every sheet says which one applies. Every numeric cell
gets a display format (D5): the decimals of its unit in `config/arrondis.yaml`, or a
scientific format for a relative gap without unit (`UNITLESS_FORMAT`); the stored value
stays exact. A sheet « Moments sur appui (E27) » gives M_nu, M_enc, 0,65 × M_enc, M_dim,
the governing term and M_axe; « Comparaison externe » starts with the structured table
of the project (`comparaison_externe.json`) when it delivers one.
"""

from __future__ import annotations

import json
import math
from collections.abc import Iterable, Sequence
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from aster_expert.arrondis import RoundingTable, load_table
from aster_expert.livrables.unites import display
from aster_expert.projet.context import ProjectContext
from aster_expert.projet.errors import ProjectError

FONT = Font(name="Arial", size=10, color="FF000000")
BOLD = Font(name="Arial", size=10, bold=True, color="FF000000")
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FF000000")
HEADER_FILL = PatternFill(
    fill_type="solid", start_color="FFD9D9D9", end_color="FFD9D9D9"
)
"""Grey 15 % (never blue: request of MM)."""
MUTED_FILL = PatternFill(
    fill_type="solid", start_color="FFF2F2F2", end_color="FFF2F2F2"
)
"""Grey 5 %: a value « non significative »."""
WORKBOOK = "classeur_livrables.xlsx"
FIRST_SHEETS = (
    "Synthèse ingénieur",
    "Vérifications de service",
    "Torseurs",
    "Moments sur appui (E27)",
)
UNITLESS_FORMAT = "0.00E+00"
"""Display of a number without unit (relative gaps, ratios): D5, lot 49 ter."""
INTEGER_FORMAT = "0"
CONVENTION_CA = (
    "Convention de Code_Aster : N > 0 en traction ; MFY = ∫z·σxx dS, MFZ = −∫y·σxx dS "
    "(R3.08.01) ; MXX, MYY > 0 tendent la face supérieure (CA-045) ; repères en "
    "feuille Conventions."
)
CONVENTION_ING = (
    "Convention ingénieur (celle des fichiers flexion_composee/*.csv) : N > 0 en "
    "compression ; M > 0 dans le sens déclaré par la section (moment de travée : fibre "
    "inférieure tendue ; moment sur appui : fibre supérieure tendue ; section "
    "circulaire : moment résultant) ; V en valeur absolue."
)
LAST_SHEETS = (
    "Réactions et équilibre",
    "Contrôles D et E",
    "Comparaison externe",
    "Conventions",
    "Provenance",
)
SHORT = {
    "ELU_fondamental": "ELU",
    "ELU_accidentel": "ELU acc.",
    "ELS_caracteristique": "ELS car.",
    "ELS_frequent": "ELS fréq.",
    "ELS_quasi_permanent": "ELS QP",
}
COMPONENT_ORDER = (
    "N", "VY", "VZ", "MT", "MFY", "MFZ", "M_RES", "V_RES",
    "NXX", "NYY", "NXY", "MXX", "MYY", "MXY", "QX", "QY",
    "DX", "DY", "DZ",
)  # fmt: skip
COMPONENT_LABELS = {"M_RES": "M résultant", "V_RES": "V résultant"}


class Sheet:
    """A worksheet written row by row with the layout of the workbook."""

    def __init__(self, workbook: Workbook, title: str, table: RoundingTable) -> None:
        self.ws = workbook.create_sheet(title[:31])
        self.table = table
        self.row = 1
        self.widths: dict[int, int] = {}
        self.frozen = False

    def text(self, value: str, *, bold: bool = False) -> None:
        cell = self.ws.cell(row=self.row, column=1, value=value)
        cell.font = BOLD if bold else FONT
        self.row += 1

    def header(
        self, labels: Sequence[str | None], *, freeze: int = 0, start: int = 1
    ) -> None:
        """A header row; the first one written with `freeze` freezes the rows above
        and including it and the `freeze` first columns."""
        for column, label in enumerate(labels, start):
            cell = self.ws.cell(row=self.row, column=column, value=label)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if label:
                self.widths[column] = max(
                    self.widths.get(column, 0), min(len(label), 18)
                )
        if freeze and not self.frozen:
            self.ws.freeze_panes = self.ws.cell(row=self.row + 1, column=freeze + 1)
            self.frozen = True
        self.row += 1

    def merged_header(self, spans: Sequence[tuple[int, int, str]]) -> None:
        """Grey header cells spanning several columns (first, last, label)."""
        for first, last, label in spans:
            for column in range(first, last + 1):
                cell = self.ws.cell(row=self.row, column=column)
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
            cell = self.ws.cell(row=self.row, column=first, value=label)
            cell.alignment = Alignment(horizontal="center", vertical="top")
            if last > first:
                self.ws.merge_cells(
                    start_row=self.row,
                    start_column=first,
                    end_row=self.row,
                    end_column=last,
                )
        self.row += 1

    def values(
        self,
        items: Sequence[Any],
        units: Sequence[str | None] = (),
        *,
        muted: bool = False,
        bold: Iterable[int] = (),
    ) -> None:
        """One row; a number with a display unit gets the format of that unit."""
        strong = set(bold)
        for column, value in enumerate(items, 1):
            if isinstance(value, float) and math.isnan(value):
                value = None
            cell = self.ws.cell(row=self.row, column=column, value=value)
            cell.font = BOLD if column in strong else FONT
            if muted:
                cell.fill = MUTED_FILL
            unit = units[column - 1] if column - 1 < len(units) else None
            if isinstance(value, float) and unit:
                decimals = self.table.decimals(unit)
                cell.number_format = "0" if decimals == 0 else "0." + "0" * decimals
            elif isinstance(value, bool):
                pass
            elif isinstance(value, float):
                cell.number_format = UNITLESS_FORMAT
            elif isinstance(value, int):
                cell.number_format = INTEGER_FORMAT
            text = "" if value is None else str(value)
            if isinstance(value, float):
                text = f"{value:.2f}"
            self.widths[column] = max(self.widths.get(column, 0), min(len(text), 45))
        self.row += 1

    def blank(self) -> None:
        self.row += 1

    def finish(self) -> None:
        for column, width in self.widths.items():
            self.ws.column_dimensions[get_column_letter(column)].width = max(
                7, width + 2
            )


def shown(value: float | None, unit_si: str) -> tuple[float | None, str]:
    """(value in its display unit, display unit): exact conversion, no rounding."""
    target = display(unit_si)
    return (None if value is None else value * target.factor), target.unit


def _load(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def _flatten_verdicts(prefix: str, data: Any) -> Iterable[tuple[str, Any, str]]:
    """(path, value, verdict) of every dict holding a `verdict` key."""
    if isinstance(data, dict):
        if "verdict" in data and isinstance(data["verdict"], str):
            value = next(
                (
                    data[k]
                    for k in (
                        "worst_relative_gap",
                        "relative_gap",
                        "ecart_relatif",
                        "ecart_relatif_max",
                        "ecart",
                        "valeur",
                    )
                    if k in data
                ),
                None,
            )
            yield prefix, value, data["verdict"]
        for key, sub in data.items():
            if key in ("verdict", "provenance"):
                continue
            yield from _flatten_verdicts(f"{prefix} / {key}" if prefix else key, sub)


def _place_text(lieu: dict[str, Any]) -> str:
    parts = [lieu["groupe"]] if lieu.get("groupe") else []
    if lieu.get("maille_med") is not None:
        parts.append(f"maille {lieu['maille_med']}")
        if lieu.get("extremite") is not None:
            parts.append(f"extr. {lieu['extremite']}")
    if lieu.get("interpolation"):
        nodes = "-".join(str(n) for n in lieu["interpolation"]["noeuds_med"])
        parts.append(f"interpolé entre les nœuds {nodes}")
    else:
        parts.append(f"nœud {lieu['noeud_med']}")
    x, y, z = lieu["coordonnees_m"]
    parts.append(f"({x:.2f} ; {y:.2f} ; {z:.2f})")
    return ", ".join(parts)


def _converted(
    t: dict[str, Any], column: str | None
) -> tuple[float | None, str | None]:
    if not column or column not in t["torseur_si"]:
        return None, None
    value, unit = shown(t["torseur_si"][column], t["unites_si"][column])
    return value, unit


def _moment_sign(section: dict[str, Any], column: str | None) -> float:
    """Factor from the Code_Aster moment to the section's declared sense (D4)."""
    if column is None:
        return 1.0
    declared = (section.get("signes") or {}).get(column)
    return -1.0 if declared == "negative" else 1.0


def engineer(
    t: dict[str, Any], section: dict[str, Any], role: str, column: str | None
) -> tuple[float | None, str | None]:
    """A value of a torsor in the engineer's convention (D4): N > 0 in compression,
    M in the section's sense, |V|."""
    value, unit = _converted(t, column)
    if value is None:
        return None, unit
    if role == "N":
        return -value, unit
    if role == "M":
        return _moment_sign(section, column) * value, unit
    return abs(value), unit


def _sheet_summary(s: Sheet, syn: dict[str, Any]) -> None:
    states = [e["etat"] for e in syn["etats_limites"]]
    s.text(
        f"Projet {syn['projet']} — synthèse ingénieur : valeur dimensionnante de "
        "chaque section et torseur concomitant (N, M, V au même lieu, même "
        "combinaison, même position mobile), par état limite ; coques par mètre de "
        "largeur.",
        bold=True,
    )
    s.text(CONVENTION_ING, bold=True)
    s.text(
        "Les torseurs bruts (convention de Code_Aster) sont dans la feuille Torseurs ; "
        "les moments sur appui (E27) dans la feuille Moments sur appui (E27)."
    )
    absent = "; ".join(f"{a['libelle']} : {a['motif']}" for a in syn["etats_absents"])
    if absent:
        s.text(f"États limites absents : {absent}")
    s.text(
        f"Fissuration : {syn['fissuration']['libelle']} "
        f"({syn['fissuration']['source']})"
    )
    fixed = ["Famille", "Section", "Grandeur", "Unités"]
    spans = []
    column = len(fixed) + 1
    for state in states:
        spans.append((column, column + 3, SHORT[state]))
        column += 4
    spans.append((column, column + 1, "Renvois"))
    s.merged_header([(1, len(fixed), "")] + spans)
    labels = list(fixed)
    for _state in states:
        labels += ["Valeur", "N (compr. > 0)", "M", "|V|"]
    labels += ["Détail", "Remarque"]
    s.header(labels, freeze=3)
    by_key: dict[tuple[str, str, bool], dict[str, dict[str, Any]]] = {}
    order: list[tuple[str, str, bool]] = []
    for t in syn["torseurs"]:
        if not t.get("principal", True) or t.get("comparaison_seulement"):
            continue
        name = t.get("libelle_ingenieur") or t["grandeur"]
        key = (t["section"], name, t["significative"])
        if key not in by_key:
            by_key[key] = {}
            order.append(key)
        by_key[key][t["etat_limite"]] = t
    sections = {x["nom"]: x for x in syn["sections"]}
    for family in syn["familles"]:
        s.blank()
        reserve = f" — réserve : {family['reserve']}" if family.get("reserve") else ""
        s.text(
            f"{family['nom']} ({family['type']}, champ {family['champ']}, groupes "
            f"{' '.join(family['groupes'])}){reserve}",
            bold=True,
        )
        for key in order:
            section = sections[key[0]]
            if section["famille"] != family["nom"]:
                continue
            entries = by_key[key]
            any_t = next(iter(entries.values()))
            triplet = any_t["triplet"]
            row: list[Any] = [family["nom"], section["libelle"], key[1]]
            units_text: list[str] = []
            units: list[str | None] = [None, None, None, None]
            bold: list[int] = []
            values: list[Any] = []
            for state in states:
                t = entries.get(state)
                if t is None:
                    values += [None] * 4
                    units += [None] * 4
                    continue
                main, unit = engineer(t, section, t["role"], t["colonne"])
                cells: list[Any] = [main]
                cell_units: list[str | None] = [unit]
                for role in ("N", "M", "V"):
                    v, u = engineer(t, section, role, triplet.get(role))
                    cells.append(v)
                    cell_units.append(u)
                bold.append(len(row) + 2 + len(values))
                values += cells
                units += cell_units
                if not units_text:
                    units_text = [u or "" for u in cell_units]
            row.append(" ; ".join(dict.fromkeys(u for u in units_text if u)))
            row += values
            files = section.get("export") or []
            row.append(
                "Torseurs ; " + ", ".join(f"flexion_composee/{f}" for f in files)
                if files
                else "Torseurs"
            )
            remark = None if key[2] else f"non significative : {any_t['motif']}"
            if "e27" in any_t:
                remark = "; ".join(
                    f"{SHORT[st]} : {x['e27']['terme_gouvernant']} gouverne"
                    for st, x in entries.items()
                )
            elif any_t.get("signe_oppose"):
                remark = any_t.get("motif")
            elif section.get("reserve") and key[2]:
                remark = section["reserve"]
            row.append(remark)
            s.values(row, units, muted=not key[2], bold=bold)


def _sheet_service(s: Sheet, data: dict[str, Any]) -> None:
    s.text(
        "Critères de service déclarés par MM, valeurs lues dans les déplacements "
        "livrés "
        "(superposition des cas élémentaires et des positions mobiles)",
        bold=True,
    )
    if not data.get("criteres"):
        s.text(f"Aucun critère de service : {data.get('motif_absence')}")
        return
    s.header(
        [
            "Critère",
            "État limite",
            "Cas exclus",
            "Valeur (mm)",
            "Limite (mm)",
            "Taux",
            "Verdict",
            "Lieu",
            "Portée",
            "Combinaison",
            "Position mobile",
            "Source",
        ],
        freeze=1,
    )
    for c in data["criteres"]:
        for key, muted in (("gouvernant", False), ("hors_verdict", True)):
            g = c.get(key)
            if not g:
                continue
            x, y, z = g["coordonnees_m"]
            place = ", ".join(
                p
                for p in (
                    g.get("groupe"),
                    f"nœud {g['noeud_med']}",
                    f"({x:.2f} ; {y:.2f} ; {z:.2f})",
                )
                if p
            )
            verdict = (
                c.get("verdict")
                if key == "gouvernant"
                else (f"sans verdict : {g.get('note')}")
            )
            s.values(
                [
                    c["libelle"],
                    c["etat_limite_libelle"],
                    ", ".join(c["cas_exclus"]) or None,
                    g["valeur_m"] * 1.0e3,
                    c["limite_m"] * 1.0e3,
                    c.get("taux") if key == "gouvernant" else g.get("taux"),
                    verdict,
                    place,
                    g.get("portee"),
                    g["combinaison"],
                    g.get("position_mobile"),
                    c["source"],
                ],
                [None, None, None, "mm", "mm"],
                muted=muted,
            )


def _sheet_torsors(s: Sheet, syn: dict[str, Any], lines: list[dict[str, Any]]) -> None:
    s.text(
        "Torseurs complets (format large), signes et repères de Code_Aster : d'abord "
        "les torseurs dimensionnants (section × état limite × grandeur), puis, pour "
        "chaque section et chaque combinaison, le torseur où la grandeur principale "
        "est la plus grande en module. Section en T (E25) : N, MFY, VZ sont les totaux "
        "barre + bande ; *_barre : la barre seule (traçabilité).",
        bold=True,
    )
    s.text(CONVENTION_CA, bold=True)
    present = {c for t in syn["torseurs"] + lines for c in t["torseur_si"]}
    comps = [c for c in COMPONENT_ORDER if c in present]
    head = [
        "Nature",
        "Section",
        "Famille",
        "État limite",
        "Grandeur",
        "Valeur",
        "Combinaison",
        "Position mobile",
        "Lieu",
        "Unités",
        *[COMPONENT_LABELS.get(c, c) for c in comps],
        "Remarque",
    ]
    s.header(head, freeze=2)
    for nature, rows in (
        ("dimensionnant", syn["torseurs"]),
        ("par combinaison", lines),
    ):
        for t in rows:
            significant = t.get("significative", True)
            values: list[Any] = []
            units: list[str | None] = []
            for c in comps:
                if c in t["torseur_si"]:
                    v, u = shown(t["torseur_si"][c], t["unites_si"][c])
                else:
                    v, u = None, None
                values.append(v)
                units.append(u)
            unit_names = sorted({u for u in units if u})
            value, shown_unit = shown(t["valeur_si"], t["unite_si"])
            s.values(
                [
                    nature,
                    t["section"],
                    t["famille"],
                    SHORT[t["etat_limite"]],
                    t["grandeur"],
                    value,
                    t["combinaison"],
                    t["position_mobile"],
                    _place_text(t["lieu"]),
                    " ; ".join(unit_names),
                    *values,
                    None if significant else f"non significative : {t.get('motif')}",
                ],
                [None] * 5 + [shown_unit] + [None] * 4 + units,
                muted=not significant,
            )


def _sheet_e27(s: Sheet, syn: dict[str, Any], derived: dict[str, Any] | None) -> None:
    rows = [t for t in syn["torseurs"] if "e27" in t]
    s.text(
        "Moment de dimensionnement sur appui (E27, règle de MM) : M_dim = max(M_nu ; "
        "0,65 × M_encastrement), combinaison par combinaison et position par "
        "position, puis enveloppe. M_encastrement : moment d'encastrement parfait des "
        "modèles auxiliaires (charges verticales seulement : freinage et température "
        "exclus, choix que MM peut renverser) ; M_axe : moment élastique de la "
        "référence à l'axe, comparaison seulement. Clause EN 1992-1-1 §5.3.2.2(3) "
        "citée d'après MM et le superviseur (le corpus du dépôt ne contient pas "
        "l'Eurocode).",
        bold=True,
    )
    s.text(
        "Valeurs signées dans le sens d'appui de la section (M > 0 : moment sur "
        "appui) ; N (compression > 0) et |V| lus au nu dans la référence, à la ligne "
        "de M_nu, même combinaison et même position."
    )
    if not rows:
        s.text("Sans objet : aucune section ne déclare de moment sur appui (E27).")
    else:
        s.header(
            [
                "Section",
                "État limite",
                "Auxiliaire",
                "M_nu",
                "M_enc",
                "0,65 × M_enc",
                "M_dim",
                "Terme qui gouverne",
                "M_axe (comparaison)",
                "N (compr. > 0)",
                "|V|",
                "Unités",
                "Combinaison",
                "Position mobile",
                "Appui",
                "Lieu au nu",
                "Lieu dans l'auxiliaire",
            ],
            freeze=2,
        )
        sections = {x["nom"]: x for x in syn["sections"]}
        for t in rows:
            e = t["e27"]
            unit_si = t["unite_si"]
            m = [shown(e[k], unit_si) for k in ("M_nu_si", "M_enc_si")]
            fm = shown(e["fraction_M_enc_si"], unit_si)
            md = shown(e["M_dim_si"], unit_si)
            ma = shown(e["M_axe_si"], unit_si)
            section = sections[t["section"]]
            n, n_unit = engineer(t, section, "N", t["triplet"].get("N"))
            v, v_unit = engineer(t, section, "V", t["triplet"].get("V"))
            s.values(
                [
                    section["libelle"],
                    SHORT[t["etat_limite"]],
                    e["auxiliaire"],
                    m[0][0],
                    m[1][0],
                    fm[0],
                    md[0],
                    e["terme_gouvernant"],
                    ma[0],
                    n,
                    v,
                    " ; ".join(dict.fromkeys(u for u in (md[1], n_unit, v_unit) if u)),
                    t["combinaison"],
                    t["position_mobile"],
                    "(" + " ; ".join(f"{c:.2f}" for c in e["appui"]) + ")",
                    _place_text(t["lieu"]),
                    _place_text(e["lieu_auxiliaire"]),
                ],
                [
                    None,
                    None,
                    None,
                    md[1],
                    md[1],
                    md[1],
                    md[1],
                    None,
                    md[1],
                    n_unit,
                    v_unit,
                ],
                bold=[7],
            )
        s.blank()
        s.text("Par appui : enveloppe de M_dim de chaque appui", bold=True)
        s.header(
            [
                "Section",
                "État limite",
                "Appui",
                "M_nu",
                "M_enc",
                "0,65 × M_enc",
                "M_dim",
                "Terme qui gouverne",
                "M_axe",
                "Combinaison",
                "Position mobile",
            ]
        )
        for t in rows:
            unit_si = t["unite_si"]
            for a in t["e27"]["par_appui"]:
                s.values(
                    [
                        sections[t["section"]]["libelle"],
                        SHORT[t["etat_limite"]],
                        "(" + " ; ".join(f"{c:.2f}" for c in a["appui"]) + ")",
                        shown(a["M_nu"], unit_si)[0],
                        shown(a["M_enc"], unit_si)[0],
                        shown(a["fraction_M_enc"], unit_si)[0],
                        shown(a["M_dim"], unit_si)[0],
                        a["terme_gouvernant"],
                        shown(a["M_axe"], unit_si)[0],
                        a["combinaison"],
                        a["position_mobile"],
                    ],
                    [None, None, None, *[shown(0.0, unit_si)[1]] * 4, None]
                    + [shown(0.0, unit_si)[1]],
                )
    if derived:
        s.blank()
        s.text(
            "Modèles dérivés de la référence (lot 49 ter) : harnais D (Σ charges + Σ "
            "réactions + Σ ressorts par cas élémentaire ; réaction verticale de chaque "
            "position mobile égale à la référence)",
            bold=True,
        )
        s.header(
            [
                "Modèle",
                "Nature",
                "Libellé",
                "Conditions ajoutées",
                "Cas exclus",
                "Écart max (cas)",
                "Écart max FZ (positions)",
                "Verdict",
            ]
        )
        for name, m in derived.items():
            control = m.get("controle") or {}
            s.values(
                [
                    name,
                    m["nature"],
                    m["libelle"],
                    m["conditions"],
                    ", ".join(m["cas_exclus"]) or None,
                    control.get("ecart_relatif_max_cas"),
                    control.get("ecart_relatif_max_positions_FZ"),
                    control.get("verdict", "non exécuté"),
                ]
            )


def _sheet_tables(s: Sheet, data: dict[str, Any]) -> None:
    """Tables a project delivers as `{"tableaux": [{"titre", "note", "colonnes":
    [{"cle", "entete", "unite"}], "lignes": [...]}]}`."""
    for table in data.get("tableaux", []):
        s.text(table["titre"], bold=True)
        if table.get("note"):
            s.text(table["note"])
        columns = table["colonnes"]
        s.header([c["entete"] for c in columns], freeze=1)
        for line in table["lignes"]:
            s.values(
                [line.get(c["cle"]) for c in columns],
                [c.get("unite") for c in columns],
                muted=bool(line.get("_grise")),
            )
        s.blank()


def _sheet_detail(
    s: Sheet, family: dict[str, Any], members: list[dict[str, Any]]
) -> None:
    rows = [m for m in members if m["famille"] == family["nom"]]
    s.text(
        f"{family['nom']} : un élément par ligne et par état limite ; valeur, cote et "
        "combinaison de chaque grandeur (signes de Code_Aster)",
        bold=True,
    )
    if not rows:
        s.text("Sans objet : la famille n'a pas de membres déclarés.")
        return
    labels = list(rows[0]["valeurs"])
    head = ["Élément", "État limite"]
    for label in labels:
        head += [label, f"cote {label}", f"combinaison {label}"]
    s.header(head, freeze=2)
    axis = "XYZ".index(family.get("axe") or "Z") if family.get("axe") else 2
    for m in rows:
        values: list[Any] = [m["groupe"], SHORT[m["etat_limite"]]]
        units: list[str | None] = [None, None]
        for label in labels:
            v = m["valeurs"].get(label)
            if v is None:
                values += [None, None, None]
                units += [None, None, None]
                continue
            value, unit = shown(v["valeur_si"], v["unite_si"])
            combination = v["combinaison"] + (
                f" @ {v['position_mobile']}" if v.get("position_mobile") else ""
            )
            values += [value, v["coordonnees_m"][axis], combination]
            units += [unit, "m", None]
        s.values(values, units)


def build_workbook(package: Path, table: RoundingTable) -> Workbook:
    vis = package / "visualisation"
    manifest = _load(vis / "manifeste.json")
    paquet = _load(package / "paquet.json")
    wb = Workbook()
    wb.remove(wb.active)
    syn_path = vis / "synthese_ingenieur.json"
    syn = _load(syn_path) if syn_path.is_file() else None
    sheets: dict[str, Sheet] = {t: Sheet(wb, t, table) for t in FIRST_SHEETS}
    details: list[tuple[dict[str, Any], Sheet]] = []
    if syn is not None:
        for family in syn["familles"]:
            if family["type"] == "poutre" and family.get("membres"):
                details.append((family, Sheet(wb, f"Détail {family['nom']}", table)))
    for title in LAST_SHEETS:
        sheets[title] = Sheet(wb, title, table)

    if syn is None:
        for title in FIRST_SHEETS:
            sheets[title].text(
                "Sans objet : le projet ne déclare pas de synthèse ingénieur "
                "(engineering_summary)."
            )
    else:
        _sheet_summary(sheets["Synthèse ingénieur"], syn)
        _sheet_service(
            sheets["Vérifications de service"],
            _load(vis / "verifications_service.json"),
        )
        _sheet_torsors(
            sheets["Torseurs"], syn, _load(vis / "torseurs_combinaisons.json")
        )
        derived_path = vis / "modeles_derives.json"
        _sheet_e27(
            sheets["Moments sur appui (E27)"],
            syn,
            _load(derived_path) if derived_path.is_file() else None,
        )
        members = _load(vis / "detail_membres.json")
        diagram = _load(vis / "diagrammes.json")
        axes = {p["famille"]: p["axe"] for p in diagram["poutres"]}
        for family, sheet in details:
            _sheet_detail(sheet, {**family, "axe": axes.get(family["nom"])}, members)

    # Réactions et équilibre (lot 49), puis équilibre par état limite (lot 49 bis)
    s = sheets["Réactions et équilibre"]
    loads = _load(vis / "controle_charges.json")
    balance = _load(vis / "equilibre_livrables.json")
    s.text(
        "Par cas : charge appliquée (Σ des forces nodales du MED des charges), "
        "réactions et forces des ressorts livrées, bilan",
        bold=True,
    )
    s.header(
        [
            "Cas",
            "Composante",
            "Charge",
            "Réactions",
            "Ressorts",
            "Bilan",
            "Unité",
            "Écart relatif",
            "Résultante déclarée",
            "Verdict CTL-013",
        ],
        freeze=1,
    )
    for case, entry in balance["cas"].items():
        load = entry.get("charge_n") or {}
        ctl = loads["cas"].get(case, {})
        scale = max(
            (abs(v) for k, v in load.items() if k in ("FX", "FY", "FZ")), default=0
        )
        for comp in ("FX", "FY", "FZ"):
            bilan = (entry.get("bilan_n") or {}).get(comp)
            declared = (ctl.get("attendu_tel_qu_evalue_n") or {}).get(
                {"FX": "DX", "FY": "DY", "FZ": "DZ"}[comp]
            )
            s.values(
                [
                    case,
                    comp,
                    shown(load.get(comp), "N")[0],
                    shown(entry["reactions_n"][comp], "N")[0],
                    shown(entry["ressorts_n"][comp], "N")[0],
                    shown(bilan, "N")[0],
                    "kN",
                    abs(bilan) / scale if bilan is not None and scale else None,
                    shown(declared, "N")[0],
                    ctl.get("verdict"),
                ],
                [None, None, "kN", "kN", "kN", "kN", None, None, "kN"],
            )
    if balance.get("appuis_non_livres"):
        s.blank()
        s.text(
            "Appuis non livrés (motif dans projet.yaml) : "
            + ", ".join(balance["appuis_non_livres"])
        )
    by_state = vis / "equilibre_etats_limites.json"
    if by_state.is_file():
        data = _load(by_state)
        s.blank()
        s.text(data["controle"], bold=True)
        s.header(
            [
                "État limite",
                "Évaluations",
                "Écart relatif max",
                "Seuil",
                "Verdict",
                "Pire combinaison",
                "Position mobile",
            ]
        )
        for entry in data["etats"].values():
            worst = entry.get("pire") or {}
            s.values(
                [
                    entry["etat_limite"],
                    entry["evaluations"],
                    entry["ecart_relatif_max"],
                    entry["seuil"],
                    entry["verdict"],
                    worst.get("combinaison"),
                    worst.get("position_mobile"),
                ]
            )

    # Contrôles D et E
    s = sheets["Contrôles D et E"]
    dates: dict[str, tuple[str | None, str | None, str]] = {}
    for c in paquet.get("criteres", []):
        dates.setdefault(
            c["niveau"], (c.get("commit_introduction"), c.get("date"), c["document"])
        )
    s.text("Niveau atteint : " + str(paquet["niveau"]["atteint"]), bold=True)
    s.header(["Niveau", "Atteint", "Motif"], freeze=1)
    for n in paquet["niveau"]["par_niveau"]:
        s.values([n["niveau"], "oui" if n["atteint"] else "non", n["motif"]])
    s.blank()
    s.header(
        [
            "Niveau",
            "Critère",
            "Valeur",
            "Verdict",
            "Document du critère",
            "Commit",
            "Date du commit",
        ]
    )
    controls = package / "controles"
    for path in sorted(controls.glob("*.json")) if controls.is_dir() else []:
        for name, value, verdict in _flatten_verdicts(path.stem, _load(path)):
            commit, date, document = dates.get("D", (None, None, None))
            s.values(
                [
                    "D",
                    name,
                    value if isinstance(value, int | float) else None,
                    verdict,
                    document,
                    commit,
                    date,
                ]
            )
    level_e = paquet.get("niveau_e") or {}
    for name, verdict in (level_e.get("verdicts") or {}).items():
        commit, date, document = dates.get("E", (None, None, None))
        s.values(["E", name, None, verdict, document, commit, date])
    for case, entry in loads["cas"].items():
        s.values(
            [
                "D",
                f"CTL-013 MED des charges / {case}",
                entry.get("ecart_relatif"),
                entry.get("verdict"),
                "lot 49",
                None,
                None,
            ]
        )
    cross = vis / "controles_croises.json"
    if cross.is_file():
        data = _load(cross)
        s.values(
            [
                "—",
                "Synthèse ingénieur relue par resultats lire "
                f"({data['valeurs_comparees']}"
                " valeurs)",
                data["ecart_relatif_max"],
                data["verdict"],
                "lot 49 bis",
                None,
                None,
            ]
        )

    # Comparaison externe
    s = sheets["Comparaison externe"]
    structured = vis / "comparaison_externe.json"
    if structured.is_file():
        _sheet_tables(s, _load(structured))
    external = package / "contre_calcul"
    files = sorted(external.glob("*.json")) if external.is_dir() else []
    if not files and not structured.is_file():
        s.text("Sans objet : le projet n'a pas de comparaison à un calcul externe.")
    for path in files:
        data = _load(path)
        found = data.get("lignes") if isinstance(data, dict) else None
        if not found or not isinstance(found, list) or not isinstance(found[0], dict):
            continue
        table_rows: list[dict[str, Any]] = found
        s.text(
            f"{path.name} -- contre-calcul d'un modèle non validé, jamais un niveau E",
            bold=True,
        )
        keys = [
            k for k in table_rows[0] if not isinstance(table_rows[0][k], dict | list)
        ]
        s.header(keys, freeze=1)
        for line in table_rows:
            s.values(
                [
                    line.get(k) if not isinstance(line.get(k), dict | list) else None
                    for k in keys
                ]
            )
        s.blank()

    # Conventions
    s = sheets["Conventions"]
    s.header(
        [
            "Famille",
            "Champ",
            "Groupes",
            "Composantes (unité SI → affichage)",
            "Repère et signes",
        ],
        freeze=1,
    )
    for c in manifest["champs"]:
        s.values(
            [
                c["famille"],
                c["champ"],
                " ".join(c["groupes"] + c["groupes_noeuds"]),
                ", ".join(
                    f"{x['nom']} ({x['unite_si']} → {x['unite_affichage']})"
                    for x in c["composantes"]
                ),
                c["convention"],
            ]
        )
    if syn is not None:
        s.blank()
        s.header(["État limite", "Familles de combinaisons", "Note"])
        for e in syn["etats_limites"]:
            note = "fissuration" if e["etat"] == syn["fissuration"]["etat"] else None
            s.values([e["libelle"], ", ".join(e["familles_de_combinaisons"]), note])
        for a in syn["etats_absents"]:
            s.values([a["libelle"], None, f"absent : {a['motif']}"])
        s.blank()
        s.header(["Section", "Famille", "Règle de lecture", "Justification", "Lecture"])
        for x in syn["sections"]:
            s.values(
                [
                    x["libelle"],
                    x["famille"],
                    x["regle"],
                    x["justification"],
                    "; ".join(x["lecture"]) or None,
                ]
            )

    # Provenance
    s = sheets["Provenance"]
    s.header(["Élément", "Valeur"], freeze=1)
    for label, value in (
        ("Projet", manifest["projet"]),
        ("Code_Aster", manifest["code_aster"].get("version")),
        ("Image", manifest["code_aster"].get("image_digest")),
        ("Exécution de production", manifest["execution_production"]),
        (".comm de production (SHA-256)", manifest["comm_production_sha256"]),
        ("Exécution de post-traitement", manifest["execution_post_traitement"]),
        (".comm de post-traitement (SHA-256)", manifest["comm_post_traitement_sha256"]),
        ("Maillage", manifest["maillage"]["chemin"]),
        ("Maillage (SHA-256)", manifest["maillage"]["sha256"]),
        ("Paquet : commit", paquet["depot"]["commit"]),
        ("Paquet : exporté le", paquet["exporte_le"]),
        (
            "Synthèse ingénieur : méthode",
            syn["methode"] if syn is not None else None,
        ),
    ):
        s.values([label, value])
    s.blank()
    s.header(["Fichier", "SHA-256", "Octets"])
    for name, info in manifest["fichiers"].items():
        s.values([name, info["sha256"], info["octets"]])
    for sheet in [*sheets.values(), *(d[1] for d in details)]:
        sheet.finish()
    return wb


def write_workbook(ctx: ProjectContext) -> str:
    """`resultats/livrables/classeur_livrables.xlsx` from the exported package."""
    section = ctx.declaration.results_package
    if section is None:
        message = "classeur : le projet ne déclare pas de paquet de résultats"
        raise ProjectError(message)
    package = ctx.results_dir / section.output
    if not (package / "visualisation" / "manifeste.json").is_file():
        message = (
            "classeur : le paquet ne porte pas la section visualisation (exporter "
            "le paquet après les livrables)"
        )
        raise ProjectError(message)
    wb = build_workbook(package, load_table(ctx.repo_root))
    out = ctx.results_dir / "livrables" / WORKBOOK
    wb.save(out)
    return f"Classeur : {out.name} ({len(wb.sheetnames)} feuilles, depuis le paquet)"


__all__ = [
    "FIRST_SHEETS",
    "FONT",
    "HEADER_FILL",
    "LAST_SHEETS",
    "MUTED_FILL",
    "UNITLESS_FORMAT",
    "WORKBOOK",
    "build_workbook",
    "engineer",
    "shown",
    "write_workbook",
]
