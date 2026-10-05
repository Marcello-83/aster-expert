"""Marqueurs `[À CONFIRMER MM]` : texte en vigueur contre historique (lot 46)."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.document import Document as DocumentType

from aster_expert.agent.hypotheses import HypothesesSheet, PracticeHypothesisEntry
from aster_expert.notes.controle import (
    ANNEX_TITLE,
    NoteCheck,
    check_markers,
    check_sheet_markers,
)


def _note(
    en_vigueur: list[str],
    historique: list[str],
    *,
    cellule: str | None = None,
    annexe: bool = True,
) -> DocumentType:
    document = Document()
    document.add_heading("1. Hypothèses", level=1)
    for text in en_vigueur:
        document.add_paragraph(text)
    if annexe:
        document.add_heading(f"2. {ANNEX_TITLE}", level=1)
        for text in historique:
            document.add_paragraph(text)
        if cellule is not None:
            table = document.add_table(rows=1, cols=1)
            table.cell(0, 0).text = cellule
    return document


def _verdict(document: DocumentType) -> NoteCheck:
    result = NoteCheck(note="n.docx")
    check_markers(document, result)
    return result


def test_note_sans_marqueur_passe() -> None:
    result = _verdict(_note(["Es = 200 GPa."], ["Décision remplacée."]))
    assert result.ok


def test_marqueur_cite_dans_l_historique_est_autorise() -> None:
    result = _verdict(
        _note(
            ["Es = 200 GPa."],
            ["Texte remplacé : « À CONFIRMER MM » (révision levée)."],
            cellule="Décision remplacée : « À CONFIRMER MM ».",
        )
    )
    assert result.ok, result.problems
    assert result.counts["marqueurs_cites_historique"] == 2


def test_marqueur_dans_un_texte_en_vigueur_echoue_cite_ou_non() -> None:
    brut = _verdict(_note(["Valeur [À CONFIRMER MM]."], []))
    assert not brut.ok
    assert "dans le texte en vigueur" in brut.problems[0]
    cite = _verdict(_note(["Valeur « À CONFIRMER MM »."], []))
    assert not cite.ok
    assert "texte en vigueur" in cite.problems[0]


def test_marqueur_ouvert_dans_l_historique_echoue() -> None:
    result = _verdict(_note(["Es = 200 GPa."], ["Ancien texte [À CONFIRMER MM]."]))
    assert not result.ok
    assert "dans le historique" in result.problems[0]


def test_marqueur_en_vigueur_dans_une_cellule_de_tableau_echoue() -> None:
    document = Document()
    document.add_heading("1. Résultats", level=1)
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "« QUESTION MM »"
    result = _verdict(document)
    assert not result.ok


def test_todo_echoue_partout() -> None:
    assert not _verdict(_note(["TODO"], [])).ok
    assert not _verdict(_note(["x"], ["TODO"])).ok


def test_un_marqueur_cite_sans_annexe_echoue() -> None:
    result = _verdict(_note(["« À CONFIRMER MM »"], [], annexe=False))
    assert not result.ok


def _sheet(option: str) -> HypothesesSheet:
    """The slab's fiche, one decision in force given a marker."""
    path = (
        Path(__file__).resolve().parents[2]
        / "projets"
        / "dalle_appuyee_01"
        / "fiche"
        / "hypotheses_dalle_arbitree.json"
    )
    sheet = HypothesesSheet.model_validate_json(path.read_text(encoding="utf-8"))
    entries = list(sheet.entries)
    for index, entry in enumerate(entries):
        if isinstance(entry, PracticeHypothesisEntry) and entry.status == "tranche":
            entries[index] = entry.model_copy(update={"option_retained": option})
            break
    return sheet.model_copy(update={"entries": tuple(entries)})


def test_marqueur_dans_la_fiche_en_vigueur_echoue() -> None:
    result = NoteCheck(note="n.docx")
    check_sheet_markers(_sheet("Valeur [À CONFIRMER MM]"), result)
    assert any("texte en vigueur" in p for p in result.problems)


def test_fiche_sans_marqueur_en_vigueur_passe() -> None:
    result = NoteCheck(note="n.docx")
    check_sheet_markers(_sheet("Valeur ferme"), result)
    assert result.ok
