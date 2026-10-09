"""Synthèse ingénieur (lot 49 bis) sur des livrables fabriqués : schéma de la section
`engineering_summary`, règles de lecture des sections, torseurs concomitants (recalculés
ici par force brute), valeurs non significatives, export pour le logiciel de flexion
composée (conventions converties), critères de service, grille des flèches de charges,
contrôle croisé par `resultats lire`, équilibre par état limite, classeur relu.

Sans Docker ni Salomé : le modèle a 9 nœuds -- un pieu vertical de deux éléments (tête
au nœud 0), un chevêtre de deux éléments selon Y par la tête, une dalle de deux mailles
qui partage deux nœuds (3 et 4) avec le chevêtre ; cas G et Q, positions P1 et P2 d'une
charge mobile."""

from __future__ import annotations

import copy
import csv
import io
import itertools
import json
import math
import shutil
from collections.abc import Callable
from pathlib import Path
from typing import Any

import numpy as np
import pytest
from openpyxl import load_workbook
from pydantic import ValidationError

from aster_expert.arrondis import load_table
from aster_expert.livrables import calcul
from aster_expert.livrables.classeur import FIRST_SHEETS, LAST_SHEETS, build_workbook
from aster_expert.livrables.donnees import Deliverables
from aster_expert.livrables.flexion import moment_sign, number
from aster_expert.livrables.ingenieur import (
    cross_check,
    equilibrium_by_state,
    step_ingenieur,
)
from aster_expert.livrables.paraview import load_grid
from aster_expert.livrables.sections import load_family, rule_mask
from aster_expert.livrables.torseurs import envelope, section_torsors
from aster_expert.projet.context import ProjectContext, load_project
from aster_expert.projet.declaration import (
    BendingPlane,
    EngineeringSummarySection,
    SectionRule,
)
from aster_expert.projet.errors import ProjectError

ROOT = Path(__file__).resolve().parents[2]

COORDS = np.array(
    [
        [0.0, 0.0, 0.0],  # 0 tête du pieu, milieu du chevêtre
        [0.0, 0.0, -1.0],
        [0.0, 0.0, -2.0],  # 2 pied (appui)
        [0.0, -1.0, 0.0],  # 3 chevêtre et dalle
        [0.0, 1.0, 0.0],  # 4 chevêtre et dalle
        [1.0, -1.0, 0.0],
        [1.0, 1.0, 0.0],
        [2.0, -1.0, 0.0],
        [2.0, 1.0, 0.0],
    ]
)
BEAMS = np.array([[0, 1], [1, 2], [3, 0], [0, 4]])
QUADS = np.array([[3, 5, 6, 4], [5, 7, 8, 6]])
BEAM_COMPS = ["N", "VY", "VZ", "MT", "MFY", "MFZ"]
SHELL_COMPS = ["NXX", "NYY", "NXY", "MXX", "MYY", "MXY", "QX", "QY"]
CASES = ("G", "Q")
POSITIONS = ("P1", "P2")
COMBINATIONS: list[dict[str, Any]] = [
    {"name": "C_ELU", "family": "ELU", "fixed": {"G": 1.35, "Q": 1.5}},
    {
        "name": "C_ELS",
        "family": "ELS",
        "fixed": {"G": 1.0},
        "mobile_coefficient": 1.0,
        "mobile_cases": ["P1", "P2"],
        "mobile_label": "essieu",
    },
    {"name": "C_ACC", "family": "ACC", "fixed": {"G": 1.0, "Q": 0.5}},
]

FIELDS: list[dict[str, Any]] = [
    {
        "family": "deplacements",
        "field": "DEPL",
        "components": ["DX", "DY", "DZ"],
        "convention": "repère global du modèle factice",
    },
    {
        "family": "efforts_coques",
        "field": "EFGE_NOEU",
        "groups": ["DALLE"],
        "components": SHELL_COMPS,
        "convention": "dalle factice, MXX > 0 tend la face supérieure",
    },
    {
        "family": "efforts_poutres",
        "field": "EFGE_ELNO",
        "groups": ["PIEUX", "CHEV"],
        "components": BEAM_COMPS,
        "convention": "repère local de chaque poutre, N > 0 traction",
    },
    {
        "family": "reactions",
        "field": "REAC_NODA",
        "node_groups": ["APPUI"],
        "components": ["DX", "DY", "DZ"],
        "convention": "repère global, appui au pied",
    },
]
DELIVERABLES: dict[str, Any] = {
    "fields": FIELDS,
    "not_relevant": [
        {
            "item": "contraintes_volumes",
            "reason": "le modèle factice n'a pas de volume",
        },
        {"item": "ressorts", "reason": "le modèle factice n'a pas de ressort"},
    ],
    "combinations": "inutilisé:par_les_tests",
    "moving_cases": ["P*"],
    "exclusion_disks": [
        {
            "node_group": "TETES",
            "radius": 0.5,
            "groups": ["DALLE"],
            "source": "disque factice des tests",
        }
    ],
}
EXPORT_PILE: dict[str, Any] = {
    "section_type": "circulaire saisie en carré équivalent",
    "h": {"value": 0.52, "unit": "m", "source": "essai"},
    "b": {"value": 0.52, "unit": "m", "source": "essai"},
    "ds": {"value": 90, "unit": "mm", "source": "essai en mm"},
    "di": {"value": 0.09, "unit": "m", "source": "essai"},
    "concrete": {"value": "C35/45", "source": "essai"},
    "steel": {"value": "B500B", "source": "essai"},
    "exposure": {"value": "XS3", "source": "essai"},
    "cnom": {"value": 0.055, "unit": "m", "source": "essai en m"},
    "equivalence": "carré de même module",
    "planes": [{"moment": "resultante", "label": "moment résultant"}],
}
EXPORT_CAP: dict[str, Any] = {
    **EXPORT_PILE,
    "section_type": "rectangulaire",
    "equivalence": None,
    "planes": [{"moment": "MFY", "top_fibre": "+z_l", "label": "flexion verticale"}],
}
EXPORT_SLAB: dict[str, Any] = {
    **EXPORT_CAP,
    "per_metre": True,
    "planes": [
        {"moment": "MXX", "top_fibre": "+n", "label": "selon X"},
        {"moment": "MYY", "top_fibre": "+n", "label": "selon Y"},
    ],
}
SUMMARY: dict[str, Any] = {
    "limit_states": [
        {"state": "ELU_fondamental", "families": ["ELU"]},
        {"state": "ELU_accidentel", "families": ["ACC"]},
        {"state": "ELS_caracteristique", "families": ["ELS"]},
    ],
    "absent_limit_states": [
        {
            "state": "ELS_frequent",
            "reason": "aucune combinaison fréquente dans l'essai",
        },
        {
            "state": "ELS_quasi_permanent",
            "reason": "aucune combinaison quasi permanente dans l'essai",
        },
    ],
    "cracking": {"state": "ELS_caracteristique", "source": "choix de l'essai factice"},
    "families": [
        {
            "name": "pieux",
            "kind": "poutre",
            "field": "EFGE_ELNO",
            "groups": ["PIEUX"],
            "members": "PIEU_*",
            "quantities": BEAM_COMPS,
            "circular": True,
            "moments": ["MFY", "MFZ"],
            "shears": ["VY", "VZ"],
            "axis": "Z",
        },
        {
            "name": "chevêtres",
            "kind": "poutre",
            "field": "EFGE_ELNO",
            "groups": ["CHEV"],
            "members": "CHEV_*",
            "quantities": BEAM_COMPS,
            "moments": ["MFY"],
            "shears": ["VZ"],
            "axis": "Y",
        },
        {
            "name": "tablier",
            "kind": "coque",
            "field": "EFGE_NOEU",
            "groups": ["DALLE"],
            "quantities": SHELL_COMPS,
        },
        {
            "name": "appuis",
            "kind": "appui",
            "field": "REAC_NODA",
            "groups": ["APPUI"],
            "quantities": ["DZ"],
        },
    ],
    "sections": [
        {
            "name": "tete",
            "label": "tête de pieu",
            "family": "pieux",
            "rule": {
                "description": "extrémités au nœud de tête",
                "node_group": "TETES",
            },
            "justification": "moment maximal du pieu factice",
            "bending_export": EXPORT_PILE,
        },
        {
            "name": "pieu",
            "label": "pieu courant",
            "family": "pieux",
            "rule": {"description": "sous la tête", "box": {"z": [-2.0, -0.5]}},
            "justification": "partie basse du pieu factice",
            "not_exported": "tube acier du modèle factice, hors béton armé",
        },
        {
            "name": "chevetre",
            "label": "chevêtre",
            "family": "chevêtres",
            "rule": {
                "description": "station la plus proche de |Y| = 0,9",
                "station": {"axis": "Y", "value": 0.9, "absolute": True},
            },
            "justification": "nu de l'appui du chevêtre factice",
            "bending_export": EXPORT_CAP,
        },
        {
            "name": "dalle",
            "label": "dalle",
            "family": "tablier",
            "rule": {"description": "toute la dalle factice"},
            "justification": "toute la dalle du modèle factice",
            "shared_with": ["CHEV"],
            "bending_export": EXPORT_SLAB,
        },
        {
            "name": "appui",
            "label": "pied",
            "family": "appuis",
            "rule": {"description": "nœud du pied", "node_group": "APPUI"},
            "justification": "réaction verticale du pied",
            "not_exported": "réaction d'appui sans section de béton armé",
        },
    ],
    "service_criteria": [
        {
            "name": "fleche",
            "label": "flèche du chevêtre",
            "kind": "relative_deflection",
            "components": ["DZ"],
            "groups": ["CHEV"],
            "limit_state": "ELS_caracteristique",
            "exclude_cases": ["G"],
            "span_m": 2.0,
            "ratio": 400,
            "spans": {"axis": "Y", "line_axis": "X", "supports": [-1.0, 1.0]},
            "source": "critère factice des tests",
        },
        {
            "name": "tete_dep",
            "label": "déplacement en tête",
            "kind": "displacement",
            "components": ["DX", "DY"],
            "node_group": "TETES",
            "limit_state": "ELS_caracteristique",
            "span_m": 2.0,
            "ratio": 400,
            "source": "critère factice des tests",
        },
    ],
    "cuts": [
        {
            "name": "coupe",
            "family": "tablier",
            "axis": "Y",
            "value": -1.0,
            "along": "X",
            "components": ["MXX"],
            "shared_with": ["CHEV"],
        }
    ],
}


def _arrays(rows: int, comps: int, seed: int) -> dict[str, np.ndarray[Any, Any]]:
    rng = np.random.default_rng(seed)
    return {
        name: rng.uniform(-1000.0, 1000.0, size=(rows, comps))
        for name in (*CASES, *POSITIONS)
    }


def _write(root: Path) -> dict[str, dict[str, np.ndarray[Any, Any]]]:
    """Livrables du modèle factice ; rend les tableaux par champ."""
    data = root / "donnees"
    data.mkdir(parents=True)
    (root / "post").mkdir()
    (root / "paquet").mkdir()
    mesh: dict[str, Any] = {
        "coords": COORDS,
        "conn_0": QUADS,
        "type_0": np.full(2, 4),
        "conn_-1": BEAMS,
        "type_-1": np.ones(4, int),
        "gma_-1_PIEUX": np.array([0, 1]),
        "gma_-1_PIEU_A": np.array([0, 1]),
        "gma_-1_CHEV": np.array([2, 3]),
        "gma_-1_CHEV_1": np.array([2, 3]),
        "gma_0_DALLE": np.array([0, 1]),
        "gno_TETES": np.array([0]),
        "gno_APPUI": np.array([2]),
    }
    np.savez(data / "maillage.npz", **mesh)
    elno_kind = np.ones(8, int)
    layouts = {
        "DEPL": (np.zeros(9, int), np.arange(9), np.zeros(9, int), np.ones(9, int), 3),
        "EFGE_NOEU": (
            np.zeros(6, int),
            np.arange(3, 9),
            np.zeros(6, int),
            np.ones(6, int),
            8,
        ),
        "EFGE_ELNO": (
            elno_kind,
            np.repeat(np.arange(4), 2),
            np.tile([0, 1], 4),
            np.full(8, -1),
            6,
        ),
        "REAC_NODA": (
            np.zeros(1, int),
            np.array([2]),
            np.zeros(1, int),
            np.ones(1, int),
            3,
        ),
    }
    comps = {
        "DEPL": ["DX", "DY", "DZ"],
        "EFGE_NOEU": SHELL_COMPS,
        "EFGE_ELNO": BEAM_COMPS,
        "REAC_NODA": ["DX", "DY", "DZ"],
    }
    out: dict[str, dict[str, np.ndarray[Any, Any]]] = {}
    for seed, (key, (kind, entity, local, level, ncomp)) in enumerate(layouts.items()):
        arrays = _arrays(len(kind), ncomp, seed)
        if key == "EFGE_NOEU":
            # Nœud 3, partagé avec le chevêtre : MXX très grand (singulier).
            for name in arrays:
                arrays[name][0, 3] = 1.0e7
        out[key] = arrays
        npz: dict[str, Any] = {
            "row_kind": kind,
            "row_entity": entity,
            "row_local": local,
            "row_level": level,
        }
        for name in CASES:
            npz[f"cas/{name}"] = arrays[name]
        for name in POSITIONS:
            npz[f"pos/{name}"] = arrays[name]
        for i, c in enumerate(COMBINATIONS):
            env = calcul.with_moving(
                calcul.combine(arrays, c["fixed"]),
                arrays,
                float(c.get("mobile_coefficient", 0.0)),
                c.get("mobile_cases", []),
            )
            npz[f"comb/{i}/max"], npz[f"comb/{i}/min"] = env.maximum, env.minimum
        np.savez(data / f"{key}.npz", **npz)
    loads: dict[str, Any] = {
        "val/FN_G": np.zeros((9, 3)),
        "row_kind/FN_G": np.zeros(9, int),
        "row_entity/FN_G": np.arange(9),
    }
    np.savez(data / "charges.npz", **loads)
    plan = {
        "resultats": [
            {
                "champs": [
                    {"cle": "DEPL", "groupes": [], "groupes_noeuds": []},
                    {"cle": "EFGE_NOEU", "groupes": ["DALLE"], "groupes_noeuds": []},
                    {
                        "cle": "EFGE_ELNO",
                        "groupes": ["PIEUX", "CHEV"],
                        "groupes_noeuds": [],
                    },
                    {"cle": "REAC_NODA", "groupes": [], "groupes_noeuds": ["APPUI"]},
                ]
            }
        ]
    }
    (root / "post" / "plan.json").write_text(json.dumps(plan), "utf-8")
    spec = {
        "cases": {n: {"elementary": n in CASES} for n in (*CASES, *POSITIONS)},
        "combinations": COMBINATIONS,
        "envelopes": True,
    }
    (root / "spec.json").write_text(json.dumps(spec), "utf-8")
    reac = out["REAC_NODA"]
    balance = {
        "cas": {
            n: {
                "charge_n": dict(
                    zip(("FX", "FY", "FZ"), (-reac[n][0]).tolist(), strict=True)
                )
            }
            for n in CASES
        },
        "appuis_non_livres": [],
    }
    (root / "paquet" / "equilibre_livrables.json").write_text(
        json.dumps(balance), "utf-8"
    )
    container = {"champs": {k: {"composantes": v} for k, v in comps.items()}}
    (root / "conteneur.json").write_text(json.dumps(container), "utf-8")
    return out


@pytest.fixture
def delivered(
    fake_project: Callable[..., Any],
) -> tuple[ProjectContext, dict[str, dict[str, np.ndarray[Any, Any]]]]:
    deliverables = copy.deepcopy(DELIVERABLES)
    deliverables["combinations"] = None
    deliverables["not_relevant"].append(
        {"item": "combinaisons", "reason": "combinaisons écrites dans spec.json"}
    )
    project = fake_project(
        deliverables=deliverables, engineering_summary=copy.deepcopy(SUMMARY)
    )
    ctx = load_project(project.name, project.repo_root)
    (project.repo_root / "config").mkdir(exist_ok=True)
    shutil.copyfile(
        ROOT / "config" / "arrondis.yaml",
        project.repo_root / "config" / "arrondis.yaml",
    )
    arrays = _write(ctx.results_dir / "livrables")
    return ctx, arrays


# --- force brute ----------------------------------------------------------------------


def _all_combinations(
    arrays: dict[str, np.ndarray[Any, Any]], family: str
) -> list[tuple[str, str | None, np.ndarray[Any, Any]]]:
    out = []
    for c in COMBINATIONS:
        if c["family"] != family:
            continue
        base = sum(k * arrays[n] for n, k in c["fixed"].items())
        if c.get("mobile_cases"):
            for p in c["mobile_cases"]:
                out.append((c["name"], p, base + c["mobile_coefficient"] * arrays[p]))
        else:
            out.append((c["name"], None, base))
    return out


# --- schéma ---------------------------------------------------------------------------


def test_the_section_is_complete_or_refused() -> None:
    EngineeringSummarySection.model_validate(SUMMARY)
    silent = copy.deepcopy(SUMMARY)
    silent["absent_limit_states"] = silent["absent_limit_states"][:1]
    with pytest.raises(ValidationError, match="ni déclaré ni absent"):
        EngineeringSummarySection.model_validate(silent)
    cracking = copy.deepcopy(SUMMARY)
    cracking["cracking"] = {"state": "ELS_frequent", "source": "absent de l'essai"}
    with pytest.raises(ValidationError, match="fissuration"):
        EngineeringSummarySection.model_validate(cracking)
    no_export = copy.deepcopy(SUMMARY)
    del no_export["sections"][1]["not_exported"]
    with pytest.raises(ValidationError, match="not_exported"):
        EngineeringSummarySection.model_validate(no_export)
    beam = copy.deepcopy(SUMMARY)
    beam["families"][1]["moments"] = []
    with pytest.raises(ValidationError, match="moments"):
        EngineeringSummarySection.model_validate(beam)
    with pytest.raises(ValidationError, match="top_fibre"):
        BendingPlane.model_validate({"moment": "MFY", "label": "sans fibre"})


def test_the_three_projects_declare_their_summary() -> None:
    for name in ("quai_cw_test_01", "dalle_appuyee_01", "estacade_pieux_01"):
        ctx = load_project(name, ROOT)
        summary = ctx.declaration.engineering_summary
        assert summary is not None, name
        assert all(len(s.justification) >= 10 for s in summary.sections)
    estacade = load_project("estacade_pieux_01", ROOT).declaration.engineering_summary
    assert estacade is not None
    assert estacade.cracking.state == "ELS_frequent"  # E23, §13.7.2 de la source
    names = {s.name for s in estacade.sections}
    assert {
        "tete_de_pieu",
        "bouchon",
        "pieu_courant",
        "chevetre_travee",
        "chevetre_nu_appui",
        "dalle_chapeau",
        "dalle_travee",
        "dalle_rive",
    } <= names


# --- règles de lecture ----------------------------------------------------------------


def test_the_reading_rules_select_the_declared_rows(delivered: Any) -> None:
    ctx, _arrays = delivered
    d = Deliverables(ctx.results_dir / "livrables")
    infos = {i.key: i for i in d.field_infos(ctx.declaration.deliverables.fields)}
    summary = ctx.declaration.engineering_summary
    pile = load_family(d, infos["EFGE_ELNO"], summary.families[0])
    cap = load_family(d, infos["EFGE_ELNO"], summary.families[1])
    assert set(pile.member) == {"PIEU_A"} and set(cap.member) == {"CHEV_1"}
    head, _ = rule_mask(pile, d.mesh, summary.sections[0].rule)
    assert sorted(pile.node[head].tolist()) == [0]
    low, _ = rule_mask(pile, d.mesh, summary.sections[1].rule)
    assert sorted(pile.node[low].tolist()) == [1, 1, 2]
    station, notes = rule_mask(cap, d.mesh, summary.sections[2].rule)
    assert sorted(np.abs(cap.xyz[station, 1]).tolist()) == [1.0, 1.0]
    assert "1.000" in notes[0]
    near = SectionRule.model_validate(
        {
            "description": "à 1 m des têtes selon Y",
            "near": {"node_group": "TETES", "axis": "Y", "distance": 1.0},
        }
    )
    at, _ = rule_mask(cap, d.mesh, near)
    assert sorted(cap.node[at].tolist()) == [3, 4]


# --- torseurs concomitants ------------------------------------------------------------


def test_the_concomitant_torsor_is_the_brute_force_one(delivered: Any) -> None:
    ctx, arrays = delivered
    d = Deliverables(ctx.results_dir / "livrables")
    infos = {i.key: i for i in d.field_infos(ctx.declaration.deliverables.fields)}
    summary = ctx.declaration.engineering_summary
    pile = load_family(d, infos["EFGE_ELNO"], summary.families[0])
    from aster_expert.livrables.sections import section_rows  # noqa: PLC0415
    from aster_expert.livrables.torseurs import combinations_of_state  # noqa: PLC0415

    sec = section_rows(pile, d.mesh, summary.sections[0], ctx.declaration.deliverables)
    combos = combinations_of_state(d.combinations, summary, "ELS_caracteristique")
    env = envelope(pile, combos, "ELS_caracteristique", [sec])
    torsors = {
        t["grandeur"]: t for t in section_torsors(pile, d.combinations, env, sec)
    }
    # Force brute : lignes du pieu au nœud de tête (élément 0, extrémité 1 = ligne 0).
    best = None
    for name, position, values in _all_combinations(arrays["EFGE_ELNO"], "ELS"):
        row = values[0]
        m = math.hypot(row[4], row[5])
        if best is None or m > best[0]:
            best = (m, name, position, row)
    assert best is not None
    t = torsors["M max (résultante)"]
    assert t["valeur_si"] == pytest.approx(best[0], rel=1e-12)
    assert (t["combinaison"], t["position_mobile"]) == (best[1], best[2])
    for k, comp in enumerate(BEAM_COMPS):
        assert t["torseur_si"][comp] == pytest.approx(best[3][k], rel=1e-12)
    assert t["lieu"]["groupe"] == "PIEU_A" and t["lieu"]["noeud_med"] == 1
    n_min = min(_all_combinations(arrays["EFGE_ELNO"], "ELS"), key=lambda e: e[2][0][0])
    assert torsors["N min"]["valeur_si"] == pytest.approx(n_min[2][0][0])
    assert torsors["N min"]["torseur_si"]["MFY"] == pytest.approx(n_min[2][0][4])


def test_a_shared_node_is_left_out_and_reported_non_significant(delivered: Any) -> None:
    ctx, _arrays = delivered
    step_ingenieur(ctx, None)
    package = ctx.results_dir / "livrables" / "paquet"
    syn = json.loads((package / "synthese_ingenieur.json").read_text("utf-8"))
    mxx = [
        t
        for t in syn["torseurs"]
        if t["section"] == "dalle"
        and t["grandeur"] == "MXX max"
        and t["etat_limite"] == "ELU_fondamental"
    ]
    kept = next(t for t in mxx if t["significative"])
    flagged = next(t for t in mxx if not t["significative"])
    assert kept["lieu"]["noeud_med"] != 4  # nœud 3 (MED 4) partagé : exclu
    assert flagged["lieu"]["noeud_med"] == 4
    assert flagged["valeur_si"] > 1.0e6 > kept["valeur_si"]
    assert "CA-058" in flagged["motif"]
    checked = json.loads((package / "controles_croises.json").read_text("utf-8"))
    assert checked["verdict"] == "RÉUSSI"
    assert checked["valeurs_comparees"] > 100
    assert checked["ecart_relatif_max"] <= 1.0e-9


def test_the_cross_check_fails_on_an_altered_value(delivered: Any) -> None:
    ctx, _arrays = delivered
    step_ingenieur(ctx, None)
    package = ctx.results_dir / "livrables" / "paquet"
    syn = json.loads((package / "synthese_ingenieur.json").read_text("utf-8"))
    torsors = syn["torseurs"][:3]
    torsors[0]["torseur_si"]["N"] *= 1.001
    d = Deliverables(ctx.results_dir / "livrables", keep_fields=True)
    infos = d.field_infos(ctx.declaration.deliverables.fields)
    result = cross_check(ctx, d, infos, torsors)
    assert result["verdict"] == "ÉCHEC"
    assert result["echecs"][0]["composante"] == "N"


# --- export pour le logiciel de flexion composée --------------------------------------


def _rows(path: Path) -> list[list[str]]:
    raw = path.read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")  # UTF-8 avec BOM
    return list(csv.reader(io.StringIO(raw.decode("utf-8-sig")), delimiter=";"))


def test_the_bending_export_converts_the_conventions(delivered: Any) -> None:
    ctx, arrays = delivered
    step_ingenieur(ctx, None)
    folder = ctx.results_dir / "livrables" / "paquet" / "flexion_composee"
    assert sorted(p.name for p in folder.iterdir()) == [
        "chevetre.csv",
        "dalle_MXX.csv",
        "dalle_MYY.csv",
        "tete.csv",
    ]
    rows = _rows(folder / "chevetre.csv")
    header = {r[0]: r[1] for r in rows if len(r) > 1}
    assert header["ds (m)"] == "0,090"  # 90 mm convertis à l'export
    assert header["cnom (mm)"] == "55"  # 0,055 m
    assert "M = −MFY" in header["Convention M"]
    table = rows.index(
        next(r for r in rows if r[:4] == ["N°", "Type", "N (kN)", "M (kN.m)"])
    )
    assert rows[table][4] == ""  # colonne vide avant la traçabilité
    lines = [r for r in rows[table + 1 :] if r and r[0].isdigit()]
    types = [r[1] for r in lines]
    assert types == sorted(types, key=["ELU", "ELUA", "ELS", "ELSfiss"].index)
    for label in ("ELU", "ELUA", "ELS", "ELSfiss"):
        assert 1 <= types.count(label) <= 4
    # M max de l'ELU : max de −MFY (fibre supérieure +z_l) ; N = −N de Code_Aster.
    elu = [r for r in lines if r[1] == "ELU"]
    m_max = next(r for r in elu if "M max" in r[5])
    values = arrays["EFGE_ELNO"]
    cap_rows = [4, 5, 6, 7]
    station = [k for k in cap_rows if abs(COORDS[BEAMS[k // 2][k % 2], 1]) == 1.0]
    expected = max(
        (-v[k][4], -v[k][0])
        for _n, _p, v in _all_combinations(values, "ELU")
        for k in station
    )
    assert m_max[3] == number(expected[0] * 1.0e-3)
    assert m_max[2] == number(expected[1] * 1.0e-3)
    assert rows[-1][0].startswith("Empreinte du MED source")
    pile = {r[0]: r[1] for r in _rows(folder / "tete.csv") if len(r) > 1}
    assert "flexion plane" in pile["Avertissement"]


def test_the_sign_of_the_moment_follows_the_declared_top_fibre() -> None:
    def plane(moment: str, fibre: str | None) -> BendingPlane:
        return BendingPlane.model_validate(
            {"moment": moment, "top_fibre": fibre, "label": "essai"}
        )

    assert moment_sign(plane("MFY", "+z_l")) == -1.0
    assert moment_sign(plane("MFY", "-z_l")) == 1.0
    assert moment_sign(plane("MFZ", "+y_l")) == 1.0
    assert moment_sign(plane("MFZ", "-y_l")) == -1.0
    assert moment_sign(plane("MXX", "+n")) == -1.0
    with pytest.raises(ProjectError, match="incompatible"):
        moment_sign(plane("MFY", "+y_l"))
    assert number(-0.04) == "0,0" and number(1234.56) == "1234,6"


# --- critères de service, équilibre, grille des flèches -------------------------------


def test_the_relative_deflection_takes_off_the_chord_of_the_supports(
    delivered: Any,
) -> None:
    ctx, arrays = delivered
    step_ingenieur(ctx, None)
    package = ctx.results_dir / "livrables" / "paquet"
    service = json.loads((package / "verifications_service.json").read_text("utf-8"))
    fleche = next(c for c in service["criteres"] if c["critere"] == "fleche")
    # Ligne du chevêtre : nœuds 3 (Y = −1), 0 (Y = 0), 4 (Y = +1) ; ELS sans G.
    best = 0.0
    for c in COMBINATIONS[1:2]:
        for p in c["mobile_cases"]:
            dz = c["mobile_coefficient"] * arrays["DEPL"][p][:, 2]
            best = max(best, abs(dz[0] - 0.5 * dz[3] - 0.5 * dz[4]))
    assert fleche["gouvernant"]["valeur_m"] == pytest.approx(best, rel=1e-12)
    assert fleche["limite_m"] == pytest.approx(2.0 / 400)
    assert fleche["verdict"] in ("RÉUSSI", "ÉCHEC")
    assert fleche["cas_exclus"] == ["G"]


def test_the_grid_of_arrows_keeps_the_sum_and_reads_as_a_distributed_load() -> None:
    x, y = np.meshgrid(np.arange(0.0, 6.0001, 0.125), np.arange(0.0, 4.0001, 0.125))
    coords = np.column_stack([x.ravel(), y.ravel(), np.zeros(x.size)])
    forces = np.zeros_like(coords)
    forces[:, 2] = -1.0
    edge = (coords[:, 0] < 1e-9) | (coords[:, 0] > 6 - 1e-9)
    forces[edge, 2] = -0.5
    points, summed, size = load_grid(coords, forces, max_arrows=200)
    assert len(points) <= 200
    assert summed.sum(axis=0) == pytest.approx(forces.sum(axis=0), rel=1e-12)
    assert size == pytest.approx(round(size / 0.125) * 0.125)
    inner = (
        (points[:, 0] > 1.0)
        & (points[:, 0] < 5.0)
        & (points[:, 1] > 1.0)
        & (points[:, 1] < 3.0)
    )
    assert np.ptp(summed[inner, 2]) == pytest.approx(0.0, abs=1e-12)


def test_the_equilibrium_by_limit_state_uses_every_position(delivered: Any) -> None:
    ctx, arrays = delivered
    d = Deliverables(ctx.results_dir / "livrables")
    infos = d.field_infos(ctx.declaration.deliverables.fields)
    package = ctx.results_dir / "livrables" / "paquet"
    reac = arrays["REAC_NODA"]
    balance = {
        "cas": {
            n: {
                "charge_n": dict(
                    zip(("FX", "FY", "FZ"), (-reac[n][0]).tolist(), strict=True)
                )
            }
            for n in CASES
        },
        "appuis_non_livres": [],
    }
    (package / "equilibre_livrables.json").write_text(json.dumps(balance), "utf-8")
    result = equilibrium_by_state(ctx, d, infos, d.combinations, package)
    elu = result["etats"]["ELU_fondamental"]
    assert elu["verdict"] == "RÉUSSI" and elu["evaluations"] == 1
    els = result["etats"]["ELS_caracteristique"]
    # Positions P1, P2 sans résultante déclarée dans controls.cases : échec dit.
    assert els["positions_sans_charge_declaree"] == ["P1", "P2"]
    assert els["verdict"] == "ÉCHEC"


# --- classeur -------------------------------------------------------------------------


def test_the_workbook_puts_the_engineering_sheets_first(delivered: Any) -> None:
    ctx, _arrays = delivered
    step_ingenieur(ctx, None)
    paquet = ctx.results_dir / "paquet"
    vis = paquet / "visualisation"
    vis.mkdir(parents=True)
    source = ctx.results_dir / "livrables" / "paquet"
    for path in source.glob("*.json"):
        shutil.copyfile(path, vis / path.name)
    (vis / "manifeste.json").write_text(
        json.dumps(
            {
                "projet": "essai",
                "code_aster": {"version": "17.4.0"},
                "execution_production": "a",
                "comm_production_sha256": "b",
                "execution_post_traitement": "c",
                "comm_post_traitement_sha256": "d",
                "maillage": {"chemin": "m", "sha256": "e"},
                "fichiers": {},
                "champs": [],
            }
        ),
        "utf-8",
    )
    (vis / "controle_charges.json").write_text(json.dumps({"cas": {}}), "utf-8")
    (vis / "equilibre_livrables.json").write_text(
        json.dumps({"cas": {}, "appuis_non_livres": []}), "utf-8"
    )
    (paquet / "paquet.json").write_text(
        json.dumps(
            {
                "depot": {"commit": "f"},
                "exporte_le": "2026-10-09",
                "niveau": {"atteint": "D", "par_niveau": []},
            }
        ),
        "utf-8",
    )
    path = ctx.results_dir / "classeur.xlsx"
    build_workbook(paquet, load_table(ROOT)).save(path)
    wb = load_workbook(path)
    assert wb.sheetnames[:3] == list(FIRST_SHEETS)
    assert wb.sheetnames[3:5] == ["Détail pieux", "Détail chevêtres"]
    assert wb.sheetnames[5:] == list(LAST_SHEETS)
    synth = wb["Synthèse ingénieur"]
    assert synth.freeze_panes is not None
    syn = json.loads((vis / "synthese_ingenieur.json").read_text("utf-8"))
    head = next(
        t
        for t in syn["torseurs"]
        if t["section"] == "tete"
        and t["grandeur"] == "M max (résultante)"
        and t["etat_limite"] == "ELU_fondamental"
    )
    row = next(
        r
        for r in synth.iter_rows()
        if r[1].value == "tête de pieu" and r[2].value == "M max (résultante)"
    )
    assert row[4].value == head["valeur_si"] * 1.0e-3  # kN·m, conversion exacte
    assert row[4].font.b
    muted = [
        r
        for r in synth.iter_rows()
        if r[-1].value and "non significative" in str(r[-1].value)
    ]
    assert muted and all(c.fill.fgColor.rgb == "FFF2F2F2" for c in muted[0][:4])
    for sheet in wb.worksheets:
        for line in sheet.iter_rows():
            for cell in line:
                if cell.value is None:
                    continue
                assert cell.font.name == "Arial" and cell.font.sz == 10
                assert cell.font.color is None or cell.font.color.rgb == "FF000000"
                fill = cell.fill.fgColor.rgb if cell.fill.fill_type else None
                assert fill in (None, "FFD9D9D9", "FFF2F2F2"), (sheet.title, fill)


# --- paquets versionnés des trois projets ---------------------------------------------


def _package(project: str) -> Path:
    return ROOT / "projets" / project / "resultats" / "paquet" / "visualisation"


@pytest.mark.parametrize(
    "project", ["estacade_pieux_01", "quai_cw_test_01", "dalle_appuyee_01"]
)
def test_the_versioned_packages_carry_a_checked_summary(project: str) -> None:
    vis = _package(project)
    checked = json.loads((vis / "controles_croises.json").read_text("utf-8"))
    assert checked["verdict"] == "RÉUSSI"
    assert checked["ecart_relatif_max"] <= 1.0e-9
    balance = json.loads((vis / "equilibre_etats_limites.json").read_text("utf-8"))
    assert all(e["verdict"] == "RÉUSSI" for e in balance["etats"].values())
    assert any((vis / "flexion_composee").glob("*.csv"))
    assert any((vis / "diagrammes").glob("*.png"))


def test_the_pile_head_of_f7_yp_finds_the_moment_of_lot_47bis() -> None:
    """Critère du § 3 de docs/essais_lot49bis.md : 435,49 kN·m (lot 47 bis, § 8 a)."""
    vis = _package("estacade_pieux_01")
    syn = json.loads((vis / "synthese_ingenieur.json").read_text("utf-8"))
    head = next(
        t
        for t in syn["torseurs"]
        if t["section"] == "tete_de_pieu"
        and t["etat_limite"] == "ELU_fondamental"
        and t["grandeur"] == "M max (résultante)"
    )
    assert abs(head["valeur_si"] * 1.0e-3 - 435.49) <= 0.005
    assert head["lieu"]["groupe"] == "PIEU_F7_YP"
    assert head["lieu"]["coordonnees_m"][2] == pytest.approx(1.30)
    assert (head["combinaison"], head["position_mobile"]) == ("111 ELU+", "TSB_P25")
    n_kn = head["torseur_si"]["N"] * 1.0e-3
    v_kn = head["torseur_si"]["V_RES"] * 1.0e-3
    rows = _rows(vis / "flexion_composee" / "tete_de_pieu.csv")
    elu = next(r for r in rows if len(r) > 5 and r[1] == "ELU" and "M max" in r[5])
    assert elu[3] == "435,5"
    assert elu[2] == number(-n_kn)  # N > 0 en compression
    assert float(elu[10].replace(",", ".")) == pytest.approx(v_kn, abs=0.05)
    assert list(itertools.islice((r[0] for r in rows), 1)) == [
        "Section de dimensionnement"
    ]
