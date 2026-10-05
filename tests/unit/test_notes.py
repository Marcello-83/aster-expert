"""Lot 45 : notes Word génériques (hypothèses, résultats) et leur contrôle automatique.

Les notes des deux projets versionnés sont construites puis contrôlées ; chaque
contrôle est aussi éprouvé en altérant le `.docx` (valeur, police, couleur, fond,
titre, marqueur) : un contrôle qui ne détecte rien n'en est pas un.
"""

from __future__ import annotations

import json
from collections.abc import Callable
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from typer.testing import CliRunner

from aster_expert.cli import app
from aster_expert.notes.controle import (
    FORBIDDEN,
    NoteCheck,
    check_hypotheses_note,
    check_results_note,
)
from aster_expert.notes.docx_style import NoteDocument, neutralise_markers
from aster_expert.notes.formats import (
    displayed_tolerance,
    fr_number,
    fr_significant,
    parse_number,
)
from aster_expert.notes.hypotheses import build_hypotheses_note, point_key
from aster_expert.notes.resultats import Package, build_results_note
from aster_expert.notes.tables import resolve_table, table_note
from aster_expert.projet.context import ProjectContext, load_project
from aster_expert.projet.declaration import NoteColumn, NoteTable
from aster_expert.projet.errors import ProjectError

_ROOT = Path(__file__).resolve().parents[2]


@pytest.fixture(scope="module")
def dalle() -> ProjectContext:
    return load_project("dalle_appuyee_01", _ROOT)


@pytest.fixture(scope="module")
def dalle_notes(
    dalle: ProjectContext, tmp_path_factory: pytest.TempPathFactory
) -> dict[str, Path]:
    folder = tmp_path_factory.mktemp("notes")
    return {
        "hypotheses": build_hypotheses_note(dalle, folder / "hyp.docx"),
        "resultats": build_results_note(dalle, folder / "res.docx"),
    }


def _headings(document: object, style: str | None = None) -> list[str]:
    """Texts of the headings of a document (`style`: `Heading 1`, ...)."""
    found = []
    for paragraph in document.paragraphs:  # type: ignore[attr-defined]
        name = paragraph.style.name if paragraph.style is not None else ""
        if name == style or (style is None and name.startswith("Heading")):
            found.append(paragraph.text)
    return found


def _problems(check: NoteCheck) -> str:
    return "\n".join(check.problems)


def _altered(source: Path, target: Path, change: Callable[..., None]) -> Path:
    document = Document(str(source))
    change(document)
    document.save(str(target))
    return target


# --- formats -------------------------------------------------------------------


def test_numbers_round_trip_exactly_through_the_french_notation() -> None:
    for value in (0.1, 33000, 2.5e-7, 1.0e11, -3.25, 0.30000000000000004):
        assert parse_number(fr_number(value)) == float(value)
    assert fr_number(2.5) == "2,5"
    assert fr_number(True) == "oui"
    assert fr_significant(0.000191743123, 6) == "0,000191743"
    assert fr_significant(12) == "12"


def test_the_displayed_tolerance_is_half_a_unit_of_the_last_digit() -> None:
    assert displayed_tolerance("1,5") == pytest.approx(0.05)
    assert displayed_tolerance("0,000191743") == pytest.approx(0.5e-5)
    assert displayed_tolerance("-1,50001e-06") == pytest.approx(0.5e-5)


def test_a_marker_quoted_by_a_text_is_shown_between_guillemets() -> None:
    text = "mention [À CONFIRMER MM] et [QUESTION Q1] restantes"
    assert (
        neutralise_markers(text)
        == "mention « À CONFIRMER MM » et « QUESTION Q1 » restantes"
    )
    assert neutralise_markers("rien [à signaler]") == "rien [à signaler]"


def test_point_labels_sort_numerically() -> None:
    labels = ["10", "2", "9", "1", None, "x"]
    ordered = sorted(labels, key=point_key)
    assert ordered == ["1", "2", "9", "10", "x", None]


# --- the layout wanted by MM ---------------------------------------------------


def test_the_layout_is_arial_10_black_with_grey_headers_and_numbered_titles(
    tmp_path: Path,
) -> None:
    note = NoteDocument(title="Essai", footer="Essai")
    note.title("Note", "sous-titre")
    assert note.heading("Premier", 1) == "1"
    assert note.heading("Sous", 2) == "1.1"
    assert note.heading("Autre", 2) == "1.2"
    assert note.heading("Second", 1) == "2"
    note.table(["A", "B"], [["1", "2"]], caption="Tableau")
    path = tmp_path / "essai.docx"
    note.save(path)

    document = Document(str(path))
    titles = _headings(document)
    assert titles == ["1. Premier", "1.1 Sous", "1.2 Autre", "2. Second"]
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            assert run.font.name == "Arial"
            assert run.font.size == Pt(10)
            assert run.font.color.rgb == RGBColor(0, 0, 0)
    header = document.tables[0].rows[0]
    for cell in header.cells:
        fills = [s.get(qn("w:fill")) for s in cell._tc.iter(qn("w:shd"))]
        assert fills == ["D9D9D9"]
    import zipfile  # noqa: PLC0415

    with zipfile.ZipFile(path) as archive:
        styles = archive.read("word/styles.xml").decode("utf-8")
    assert "themeColor" not in styles
    assert "Calibri" not in styles


def test_a_heading_beyond_level_3_is_refused() -> None:
    note = NoteDocument(title="x", footer="x")
    with pytest.raises(ValueError, match="niveau de titre"):
        note.heading("trop profond", 4)


# --- hypotheses note -------------------------------------------------------------


def test_the_hypotheses_note_of_a_project_passes_its_check(
    dalle: ProjectContext, dalle_notes: dict[str, Path]
) -> None:
    check = check_hypotheses_note(dalle, dalle_notes["hypotheses"])

    assert check.ok, _problems(check)
    assert check.counts["points_verifies"] == 9
    assert check.counts["valeurs_de_donnees_verifiees"] >= 38
    assert check.counts["textes_de_decision_verifies"] > 0
    assert check.counts["runs_verifies"] > 100
    assert check.counts["cellules_d_en_tete_grises"] > 0


def test_the_hypotheses_note_has_the_wanted_sections(
    dalle_notes: dict[str, Path],
) -> None:
    document = Document(str(dalle_notes["hypotheses"]))
    titles = _headings(document, "Heading 1")
    assert [t.split(". ", 1)[1] for t in titles] == [
        "En-tête",
        "Objet et périmètre",
        "Données",
        "Points de la fiche",
        "Constats Code_Aster cités",
        "Combinaisons et coefficients",
        "Annexe : historique des révisions",
    ]


def test_an_altered_data_value_is_detected(
    dalle: ProjectContext, dalle_notes: dict[str, Path], tmp_path: Path
) -> None:
    def change(document: object) -> None:
        for table in document.tables:  # type: ignore[attr-defined]
            for row in table.rows:
                if row.cells[0].text == "longueur" and row.cells[1].text == "6,0":
                    row.cells[1].paragraphs[0].runs[0].text = "6,1"

    altered = _altered(dalle_notes["hypotheses"], tmp_path / "v.docx", change)
    check = check_hypotheses_note(dalle, altered)
    assert not check.ok
    assert "longueur" in _problems(check)


def test_a_missing_point_title_is_detected(
    dalle: ProjectContext, dalle_notes: dict[str, Path], tmp_path: Path
) -> None:
    def change(document: object) -> None:
        for paragraph in document.paragraphs:  # type: ignore[attr-defined]
            if paragraph.text.endswith("Point 6"):
                paragraph.runs[0].text = "4.x Rubrique"

    altered = _altered(dalle_notes["hypotheses"], tmp_path / "p.docx", change)
    check = check_hypotheses_note(dalle, altered)
    assert "Point 6 : aucun titre" in _problems(check)


def test_a_rewritten_decision_text_is_detected(
    dalle: ProjectContext, dalle_notes: dict[str, Path], tmp_path: Path
) -> None:
    def change(document: object) -> None:
        for paragraph in document.paragraphs:  # type: ignore[attr-defined]
            if paragraph.text.startswith("Option retenue : Calcul élastique"):
                paragraph.runs[1].text = "Calcul élastique et fissuré."

    altered = _altered(dalle_notes["hypotheses"], tmp_path / "t.docx", change)
    check = check_hypotheses_note(dalle, altered)
    assert "option_retained absent mot pour mot" in _problems(check)


@pytest.mark.parametrize("marker", ["[À CONFIRMER MM]", "[QUESTION 3]", "TODO"])
def test_a_left_over_marker_is_detected(
    dalle: ProjectContext,
    dalle_notes: dict[str, Path],
    tmp_path: Path,
    marker: str,
) -> None:
    def change(document: object) -> None:
        document.paragraphs[0].runs[0].text += f" {marker}"  # type: ignore[attr-defined]

    for kind, checker in (
        ("hypotheses", check_hypotheses_note),
        ("resultats", check_results_note),
    ):
        altered = _altered(dalle_notes[kind], tmp_path / f"m_{kind}.docx", change)
        found = next(m for m in FORBIDDEN if marker.startswith(m))
        assert f"« {found} »" in _problems(checker(dalle, altered))


# --- the layout, checked back from the file -----------------------------------------


@pytest.mark.parametrize(
    ("change_name", "expected"),
    [
        ("font", "police 'Calibri'"),
        ("size", "taille"),
        ("color", "couleur"),
        ("blue_header", "fond gris"),
        ("body_fill", "fond"),
        ("unnumbered", "titre non numéroté"),
        ("not_bold_header", "hors gras"),
    ],
)
def test_each_layout_deviation_is_detected(
    dalle: ProjectContext,
    dalle_notes: dict[str, Path],
    tmp_path: Path,
    change_name: str,
    expected: str,
) -> None:
    def change(document: object) -> None:
        doc = document
        run = doc.paragraphs[2].runs[0]  # type: ignore[attr-defined]
        table = doc.tables[0]  # type: ignore[attr-defined]
        if change_name == "font":
            run.font.name = "Calibri"
        elif change_name == "size":
            run.font.size = Pt(11)
        elif change_name == "color":
            run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
        elif change_name == "blue_header":
            shading = next(table.rows[0].cells[0]._tc.iter(qn("w:shd")))
            shading.set(qn("w:fill"), "4472C4")
        elif change_name == "body_fill":
            properties = table.rows[1].cells[0]._tc.get_or_add_tcPr()
            shading = OxmlElement("w:shd")
            shading.set(qn("w:val"), "clear")
            shading.set(qn("w:fill"), "FFFF00")
            properties.append(shading)
        elif change_name == "unnumbered":
            for paragraph in doc.paragraphs:  # type: ignore[attr-defined]
                if paragraph.text.startswith("1. "):
                    paragraph.runs[0].text = "Sans numéro"
                    break
        elif change_name == "not_bold_header":
            table.rows[0].cells[0].paragraphs[0].runs[0].font.bold = False

    altered = _altered(
        dalle_notes["hypotheses"], tmp_path / f"{change_name}.docx", change
    )
    check = check_hypotheses_note(dalle, altered)
    assert not check.ok
    assert expected in _problems(check)


# --- results note -----------------------------------------------------------------


def test_the_results_note_passes_its_check_and_states_the_level(
    dalle: ProjectContext, dalle_notes: dict[str, Path]
) -> None:
    check = check_results_note(dalle, dalle_notes["resultats"])

    assert check.ok, _problems(check)
    assert check.counts["cellules_de_tableau_verifiees"] > 100
    assert check.counts["valeurs_mesurees_retrouvees_dans_le_paquet"] > 50
    document = Document(str(dalle_notes["resultats"]))
    text = "\n".join(p.text for p in document.paragraphs)
    assert "Niveau atteint : E" in text


def test_the_sections_follow_the_package_not_the_project_name(
    dalle: ProjectContext, dalle_notes: dict[str, Path]
) -> None:
    package = Package(dalle)
    assert "combinaisons" not in package.sections()
    document = Document(str(dalle_notes["resultats"]))
    titles = _headings(document, "Heading 1")
    assert not any("Combinaisons" in t for t in titles)
    assert any("Ferraillage" in t for t in titles)
    assert any("Résultats par cas" in t for t in titles)
    assert any("Limites" in t for t in titles)


def test_an_altered_table_cell_is_detected(
    dalle: ProjectContext, dalle_notes: dict[str, Path], tmp_path: Path
) -> None:
    def change(document: object) -> None:
        for table in document.tables:  # type: ignore[attr-defined]
            if table.rows[0].cells[0].text == "Taille de maille (m)":
                table.rows[1].cells[2].paragraphs[0].runs[0].text = "0,5"
                return

    altered = _altered(dalle_notes["resultats"], tmp_path / "c.docx", change)
    check = check_results_note(dalle, altered)
    assert not check.ok
    assert "Cas G, suite de convergence" in _problems(check)


def test_a_measured_value_absent_from_the_package_is_detected(
    dalle: ProjectContext, dalle_notes: dict[str, Path], tmp_path: Path
) -> None:
    def change(document: object) -> None:
        for table in document.tables:  # type: ignore[attr-defined]
            if (
                table.rows[0].cells[0].text == "Grandeur"
                and table.rows[1].cells[0].text
            ):
                table.rows[1].cells[1].paragraphs[0].runs[0].text = "4,56789e-05"
                return

    altered = _altered(dalle_notes["resultats"], tmp_path / "m.docx", change)
    check = check_results_note(dalle, altered)
    assert "introuvable dans les fichiers du paquet" in _problems(check)


def test_the_results_note_states_the_open_anomaly_on_the_reinforcement(
    dalle_notes: dict[str, Path],
) -> None:
    document = Document(str(dalle_notes["resultats"]))
    text = "\n".join(c.text for t in document.tables for r in t.rows for c in r.cells)
    assert "ANO-001" in text


def test_the_results_note_has_no_verdict_of_conformity(
    dalle_notes: dict[str, Path],
) -> None:
    document = Document(str(dalle_notes["resultats"]))
    text = "\n".join(p.text for p in document.paragraphs).lower()
    for word in ("conforme", "conformité", "non conforme"):
        assert word not in text


# --- declared tables ---------------------------------------------------------------


def _table(**changes: object) -> NoteTable:
    spec: dict[str, object] = {
        "section": "cas",
        "title": "Essai",
        "source": "d.json",
        "rows": "cases.*.ties",
        "columns": [
            {"header": "Cas", "path": "^$key"},
            {"header": "Clé", "path": "$key"},
            {"header": "Effort", "path": "n"},
            {"header": "Hauteur", "path": "^h"},
        ],
    }
    spec.update(changes)
    return NoteTable.model_validate(spec)


def test_a_wildcard_path_reads_the_enclosing_element_and_the_keys(
    tmp_path: Path,
) -> None:
    data = {
        "cases": {
            "G": {"h": 2.5, "ties": [{"n": 1.0}, {"n": 2.5}]},
            "Q": {"h": 3.5, "ties": {"A": {"n": 0.000123456789}}},
        }
    }
    (tmp_path / "d.json").write_text(json.dumps(data), encoding="utf-8")

    resolved = resolve_table(tmp_path, _table())

    assert resolved.headers == ["Cas", "Clé", "Effort", "Hauteur"]
    assert resolved.values == [
        ["G", "0", 1.0, 2.5],
        ["G", "1", 2.5, 2.5],
        ["Q", "A", 0.000123456789, 3.5],
    ]
    assert resolved.cells[2] == ["Q", "A", "0,000123457", "3,5"]
    assert table_note(resolved) is None


def test_a_csv_table_keeps_one_row_in_n_and_says_so(tmp_path: Path) -> None:
    lines = ["x,y"] + [f"{i},{i * 0.5}" for i in range(10)]
    (tmp_path / "t.csv").write_text("\n".join(lines) + "\n", encoding="utf-8")
    table = NoteTable(
        section="cas",
        title="Essai",
        source="t.csv",
        columns=(
            NoteColumn(header="x", path="x"),
            NoteColumn(header="y", path="y", format="{:.2f}"),
        ),
        step=4,
    )

    resolved = resolve_table(tmp_path, table)

    assert [row[0] for row in resolved.values] == [0.0, 4.0, 8.0]
    assert resolved.cells[1] == ["4", "2,00"]
    note = table_note(resolved)
    assert note is not None
    assert "3 lignes sur 10" in note
    assert "une ligne sur 4" in note


def test_a_table_refuses_a_missing_file_a_missing_key_and_a_missing_column(
    tmp_path: Path,
) -> None:
    with pytest.raises(ProjectError, match="fichier absent"):
        resolve_table(tmp_path, _table())
    (tmp_path / "d.json").write_text(json.dumps({"cases": {"G": {}}}), encoding="utf-8")
    with pytest.raises(ProjectError, match="clé « ties » absente"):
        resolve_table(tmp_path, _table())
    (tmp_path / "t.csv").write_text("x\n1\n", encoding="utf-8")
    bad = NoteTable(
        section="cas",
        title="E",
        source="t.csv",
        columns=(NoteColumn(header="z", path="z"),),
    )
    with pytest.raises(ProjectError, match="colonne « z » absente"):
        resolve_table(tmp_path, bad)
    shallow = _table(columns=[{"header": "H", "path": "^^h"}])
    (tmp_path / "d.json").write_text(
        json.dumps({"cases": {"G": {"ties": [{"n": 1}]}}}), encoding="utf-8"
    )
    with pytest.raises(ProjectError, match="remonte plus haut"):
        resolve_table(tmp_path, shallow)


# --- command line -------------------------------------------------------------------


def test_the_cli_writes_and_checks_the_notes(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    monkeypatch.setenv("ASTER_EXPERT_ROOT", str(_ROOT))
    runner = CliRunner()
    target = tmp_path / "h.docx"

    done = runner.invoke(
        app, ["note", "hypotheses", "dalle_appuyee_01", "--sortie", str(target)]
    )

    assert done.exit_code == 0, done.output
    assert target.is_file()
    assert "Contrôle de h.docx : OK" in done.output
    refused = runner.invoke(app, ["note", "resultats", "projet_inconnu"])
    assert refused.exit_code == 1


def test_a_nan_in_the_package_does_not_hide_present_values(tmp_path: Path) -> None:
    # Lot 47 : un écart relatif NaN (cas de charge nulle) rompait la recherche.
    import json as _json

    from aster_expert.notes.controle import _in_package, _numbers_of_package

    (tmp_path / "a.json").write_text(
        _json.dumps({"x": [303.2129, float("nan"), 1.5, 990.336, 12.0]}),
        encoding="utf-8",
    )
    numbers = sorted(_numbers_of_package(tmp_path))
    assert all(n == n for n in numbers)
    assert _in_package(numbers, "303,213") and _in_package(numbers, "990,336")


def test_an_undefined_value_is_shown_as_nan() -> None:
    from aster_expert.notes.controle import _tolerates

    assert _tolerates("nan", float("nan"))
    assert not _tolerates("0", float("nan"))
